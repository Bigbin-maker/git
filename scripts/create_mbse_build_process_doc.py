from __future__ import annotations

from datetime import date
from pathlib import Path

from docx.shared import Pt

from create_mbse_framework_doc import (
    BORDER,
    CALLOUT,
    GRAY,
    LIGHT_BLUE,
    LIGHT_GRAY,
    NAVY,
    add_body,
    add_bullet,
    add_callout,
    add_heading,
    add_number,
    add_rule,
    add_table,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    setup_document,
    shade_cell,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "AI_MBSE_系统搭建过程说明_第一视角.docx"


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("AI 赋能 MBSE 系统搭建过程说明")
    set_run_font(r, 23, True, NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("以第一视角梳理从基础框架到 RAG、生成、人工确认、预评审校验和模型入库的建设过程")
    set_run_font(r, 13.5, False, GRAY)

    metadata = [
        ("写作视角", "第一视角：我如何一步步搭建当前 AI-MBSE-Demo 原型系统"),
        ("覆盖功能", "人员权限管理、RAG 知识库、大模型交互、SysML 生成、人为回路、预评审校验、模型入库、MagicDraw 同步"),
        ("文档用途", "项目汇报、论文实现过程、系统建设总结、后续开发交接"),
        ("版本日期", date.today().isoformat()),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1.35, 5.15])
    for idx, (label, value) in enumerate(metadata):
        shade_cell(table.cell(idx, 0), LIGHT_GRAY)
        set_cell_text(table.cell(idx, 0), label, bold=True, color=NAVY)
        set_cell_text(table.cell(idx, 1), value)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    add_rule(rule)


def add_stage(doc, title: str, goal: str, steps: list[str], result: str, verification: str | None = None):
    add_heading(doc, title, 2)
    add_body(doc, f"我的目标：{goal}", bold_prefix="我的目标：")
    add_body(doc, "我的搭建过程：", bold_prefix="我的搭建过程：")
    for step in steps:
        add_number(doc, step)
    add_body(doc, f"形成结果：{result}", bold_prefix="形成结果：")
    if verification:
        add_body(doc, f"验证方式：{verification}", bold_prefix="验证方式：")


def build_doc():
    doc = setup_document()
    add_title_block(doc)

    add_heading(doc, "1. 我搭建这个系统时的总体思路", 1)
    add_body(
        doc,
        "我不是直接从“让大模型生成一张图”开始搭建，而是先把它设计成一个可控的 MBSE 工作闭环：设计师提出问题或需求，系统通过知识库和大模型辅助理解，再生成候选 SysML 模型；候选模型必须经过人工修改、预评审校验和采纳确认，最后才进入模型库并推送到 MagicDraw。",
    )
    add_body(
        doc,
        "我始终坚持一个核心边界：AI 只能生成候选方案，不能绕过设计师直接入库。这样既保留大模型的效率，又保证系统工程建模过程可追溯、可审核、可修正。",
        bold_prefix="我始终坚持一个核心边界：",
    )
    add_callout(
        doc,
        "我的建设主线",
        "先搭交互和数据底座，再接入大模型和知识库；先让模型能生成，再让模型能被人工修改；先让草稿能被采纳入库，再加入预评审校验和一键修复。每一步都尽量形成可运行、可验证的小闭环。",
    )

    add_heading(doc, "2. 我最终形成的完整业务流程", 1)
    flow = [
        "我先让设计师进入项目工作区，选择已有项目或新建项目。",
        "设计师在 AI 智能体输入需求、问题或建模意图，也可以上传临时附件或知识库文档。",
        "系统根据对话模式决定是直接问答，还是调用 RAG 知识库进行增强问答。",
        "当设计师点击“生成”时，我把当前对话、附件摘要和知识上下文送入模型生成服务，产出 SysML 候选草稿。",
        "候选草稿显示在 SysML textual view 和模型画布中，但此时不进入正式模型库。",
        "设计师可以人工编辑元素、修改连接线、新增连接线或删除连接线。",
        "设计师点击“执行模型校验”后，系统弹出《预评审报告》，列出错误、警告、提示和可修复项。",
        "设计师可以对单项问题或全部可修复问题执行“一键修复”，系统更新草稿并重算 SysML textual view。",
        "设计师确认后点击“采纳”，我再把当前草稿写入正式模型库。",
        "采纳后的模型可以继续推送 MagicDraw、导出交换包，或进入仿真流程。",
    ]
    for item in flow:
        add_number(doc, item)

    add_heading(doc, "3. 我是如何分阶段搭建各个功能的", 1)

    add_stage(
        doc,
        "3.1 第一阶段：先搭建前后端基础框架",
        "让系统先能运行起来，前端有清晰的工作台，后端有统一 API、数据库和健康检查能力。",
        [
            "我先确定前端采用 React/Vite，后端采用 FastAPI，数据库采用 SQLite，便于快速迭代和本地演示。",
            "我搭建了系统主界面，包括左侧项目栏、中间 AI 智能体、右侧 SysML 生成区和 MagicDraw 同步区。",
            "我在后端建立统一的数据库连接、配置读取、路由注册和健康检查接口。",
            "我把前端 API 调用集中封装到 client 中，避免每个组件自己拼接口地址。",
            "我先保证 health、chat、documents、sysml、magicdraw、simulation 等路由可以被统一挂载。",
        ],
        "系统形成了一个可启动、可访问、前后端能通信的基础工作台。",
        "通过前端构建、后端编译和 /api/health 接口检查基础框架是否正常。",
    )

    add_stage(
        doc,
        "3.2 第二阶段：加入项目与人员权限管理",
        "让系统不是一个单一聊天窗口，而是能够围绕不同项目组织工作，并为后续权限控制留出接口。",
        [
            "我先把原来的“对话”概念调整为“项目”，因为 MBSE 工作更适合围绕项目组织。",
            "我在前端建立项目列表、新建项目、项目重命名、项目删除和项目管理入口。",
            "我把每个项目的聊天记录、模型草稿、推送结果、校验报告都隔离保存，避免项目之间互相污染。",
            "我加入受保护基线项目的概念，普通设计师不能随意删除系统内置项目。",
            "我把权限控制先放在前端原型中实现，后续可迁移到后端用户表和角色矩阵。",
        ],
        "系统具备项目级工作区，每个项目拥有独立的聊天、模型和校验状态。",
        "通过切换项目、新建项目、删除受保护项目等操作验证权限逻辑。",
    )

    add_stage(
        doc,
        "3.3 第三阶段：接入大模型交互",
        "让设计师可以通过自然语言和系统交互，完成需求澄清、普通问答和建模意图表达。",
        [
            "我先封装 LLMService，把模型地址、Key、模型名称和真实/Mock 模式统一管理。",
            "我实现普通 chat 接口和 stream chat 接口，使前端可以流式显示大模型回答。",
            "我把对话模式设计为快速回答、深度思考和评审模式，便于适应不同任务。",
            "我加入建模意图识别，避免用户只是咨询问题时系统误触发 SysML 生成。",
            "我把最近的对话上下文传给生成接口，让模型生成时能理解当前项目语境。",
        ],
        "系统具备了 AI 智能体能力，可以回答问题、解释设计内容，也可以识别建模任务。",
        "通过流式输出、普通问答和建模生成触发测试验证大模型交互。",
    )

    add_stage(
        doc,
        "3.4 第四阶段：构建 RAG 知识库",
        "让系统回答问题时不只依赖大模型自身知识，而是能引用用户上传的文献、规范和项目资料。",
        [
            "我先区分了两类上传：临时附件和知识库文档。临时附件只进入当前对话，知识库文档会长期入库。",
            "我建立 documents 表保存文档元数据和文本内容，建立 knowledge_chunks 表保存切片。",
            "我为 PDF 解析接入 pypdf，解决 PDF 上传后只得到一个片段的问题。",
            "我实现文档切片、关键词抽取和轻量检索，把相关片段拼接进大模型上下文。",
            "我在前端增加知识库状态条，显示文档数量、片段数量、上传入库、刷新和重置按钮。",
            "我增加手动重置知识库功能，避免测试上传的文档一直留在系统中。",
        ],
        "系统可以上传文档构建知识库，并在问答时调用知识片段生成更贴近资料的回答。",
        "通过上传卫星通信相关 PDF、查看片段数、执行知识库问答和重置知识库来验证。",
    )

    add_stage(
        doc,
        "3.5 第五阶段：实现 SysML 候选模型生成",
        "把自然语言需求和对话上下文转化为可预览的 SysML 模型，而不是只返回一段文字。",
        [
            "我定义了 ModelElement 和 ModelRelationship 两类核心结构，用来表达需求、模块、接口、用例、活动、约束和它们之间的关系。",
            "我实现 ModelGenerationService，优先使用大模型结构化输出，失败时用动态兜底或行业模板生成候选模型。",
            "我把结构化元素和关系渲染成 SysML textual view，让生成结果既能画图也能看到文本模型。",
            "我在前端把 SysML textual view 和模型画布绑定到同一份 projectResult，保证二者同步变化。",
            "我把生成结果保存为草稿，而不是直接写入正式模型库。",
        ],
        "点击“生成”后，系统会生成候选 SysML 工程，展示元素、关系、文本和画布。",
        "通过输入低轨通信卫星、宽带载荷等需求，观察生成元素数量、关系数量和 textual view 是否一致。",
    )

    add_stage(
        doc,
        "3.6 第六阶段：加入人为回路确认",
        "保证 AI 生成结果必须经过设计师确认，设计师的修改和采纳行为都能被记录下来。",
        [
            "我新增 model_drafts 表，用 original_json 保存 AI 初稿，用 current_json 保存人工修改后的当前草稿。",
            "我新增 model_feedback_logs 表，记录采纳、拒绝、元素修改、关系修改、连接线新增/删除、一键修复等操作。",
            "我在前端加入“预览与确认”区域，候选模型生成后先停留在这里。",
            "我加入元素编辑功能，可以修改名称、类型、包路径、来源需求、状态和描述。",
            "我加入连接线编辑功能，可以修改已有关系、新增关系和删除关系。",
            "我限制未采纳草稿不能直接推送 MagicDraw，避免绕过人为确认。",
        ],
        "系统形成了 AI 生成、人工修改、采纳或拒绝的闭环，所有关键操作都有反馈记录。",
        "通过修改元素、增删连接线、拒绝草稿和采纳草稿验证人为回路。",
    )

    add_stage(
        doc,
        "3.7 第七阶段：实现模型入库",
        "把“采纳”定义为正式边界：只有设计师采纳后的 current_json 才能进入模型库。",
        [
            "我实现 ModelDraftService.accept，把当前草稿中的 elements 和 relationships 作为真实入库对象。",
            "我通过 SysMLAdapter.commit_model 写入 sysml_elements 和 sysml_relationships。",
            "如果真实 SysML v2 API 可用，我优先尝试远程提交；如果不可用，就降级为本地 SQLite 保存。",
            "我在前端把采纳后的状态显示为“已写入模型库，可继续推送 MagicDraw”。",
            "我明确区分草稿库和正式模型库：model_drafts 是候选与修改过程，sysml_elements/sysml_relationships 是正式模型库。",
        ],
        "系统可以把人工确认后的模型写入正式模型库，并作为后续 MagicDraw 同步的数据来源。",
        "通过采纳草稿后查询 sysml_elements、sysml_relationships 或调用模型图接口验证。",
    )

    add_stage(
        doc,
        "3.8 第八阶段：加入模型预评审与智能化校验",
        "在模型正式入库前，让系统自动检查 SysML 语法、建模规范和逻辑一致性，并给出可执行修复建议。",
        [
            "我没有直接让大模型判断模型对错，而是先实现确定性规则引擎，保证校验结果稳定可复现。",
            "我新增 DraftValidationService，直接读取 model_drafts.current_json 执行草稿校验。",
            "我设计 DraftValidationIssue 和 DraftValidationReport，报告中包含严重级别、规则 ID、对象 ID、问题描述、建议和是否可修复。",
            "我实现了关系端点不存在、关系类型不规范、需求缺少 satisfy、接口缺少 allocate、约束缺少 verify、包路径不匹配、属性未定义类型、文本不同步等规则。",
            "我为可自动修复的问题绑定 fix_action，例如 normalize_package、add_relationship、set_attribute_type、rerender_sysml_text。",
            "我在前端加入“执行模型校验”按钮，弹出《预评审报告》，每条问题提供“一键修复”，底部提供“修复全部可修复项”。",
            "我让修复动作更新草稿 current_json、重算 SysML textual view，并写入 feedback log。",
        ],
        "系统可以在采纳前自动生成预评审报告，并对规范化问题执行一键修复。",
        "通过生成草稿、执行校验、修复若干 issue、观察草稿状态和反馈日志验证。",
    )

    add_stage(
        doc,
        "3.9 第九阶段：接入 MagicDraw 同步、导出与仿真入口",
        "让已采纳模型能够继续进入外部工程工具链，而不是停留在网页预览中。",
        [
            "我实现 MagicDraw Bridge 状态检查，并在顶部状态条显示连接状态。",
            "我把推送动作限制在采纳之后，推送时使用正式模型元素和关系生成 exchange_payload。",
            "我提供导出功能，可以把当前工程包保存为 JSON，便于离线查看或交换。",
            "我保留仿真入口，后端可根据配置调用 GMAT、Simulink 或 Mock 仿真结果。",
            "我把 MagicDraw 同步结果和仿真摘要也挂到项目工作区中，避免不同项目数据混在一起。",
        ],
        "系统从 AI 辅助建模进一步连接到专业建模工具和仿真流程。",
        "通过 MagicDraw 状态、推送按钮、导出包和仿真运行入口进行验证。",
    )

    add_heading(doc, "4. 每个核心功能包含什么", 1)
    add_table(
        doc,
        ["功能", "我实现的内容", "搭建重点"],
        [
            ["人员权限管理", "项目侧栏、新建项目、受保护项目、项目级工作区状态。", "先做项目隔离，再逐步扩展真实用户和角色权限。"],
            ["RAG 数据库构建", "知识库上传、PDF 解析、切片、检索、状态展示、重置。", "区分临时附件和知识库文档，避免所有上传都长期入库。"],
            ["大模型交互", "普通问答、流式回答、问答模式、建模意图识别。", "让 LLM 先服务需求澄清，再服务模型生成。"],
            ["SysML 生成", "元素、关系、SysML 文本、模型画布、草稿创建。", "AI 输出必须结构化，才能被修改、校验和入库。"],
            ["人为回路", "元素编辑、关系编辑、连接线增删、采纳、拒绝、反馈日志。", "入库对象必须是 current_json，而不是 AI 原始结果。"],
            ["预评审校验", "规则检查、评分、报告、单项修复、全部修复。", "确定性规则为主，大模型后续可用于解释和建议增强。"],
            ["模型入库", "采纳后写入 sysml_elements 和 sysml_relationships。", "把草稿阶段和正式模型库分开。"],
            ["外部同步", "MagicDraw Bridge、交换包导出、仿真入口。", "采纳之后再进入外部工具链。"],
        ],
        [1.25, 3.15, 2.1],
        LIGHT_BLUE,
    )

    add_heading(doc, "5. 我在搭建过程中解决的几个关键问题", 1)
    add_heading(doc, "5.1 知识库上传和临时附件容易混淆", 2)
    add_body(doc, "一开始如果所有上传都进入知识库，会导致测试文件长期存在，影响后续问答。我后来把上传分成两条路径：纸夹上传是临时附件，只进当前对话；知识库上传才持久化到 documents 和 knowledge_chunks。")

    add_heading(doc, "5.2 PDF 只有一个片段的问题", 2)
    add_body(doc, "上传 PDF 后曾出现“1 文档 / 1 片段”的情况。原因是没有真正解析 PDF 正文，只保存了占位文本。我接入 pypdf 后重新解析文献，切片数量恢复正常。")

    add_heading(doc, "5.3 AI 生成结果不能直接入库", 2)
    add_body(doc, "为了避免 AI 生成错误模型直接污染模型库，我新增草稿机制。生成结果进入 model_drafts，设计师修改和校验都作用于 current_json；只有采纳时才写入正式模型库。")

    add_heading(doc, "5.4 人工修改必须真的改变模型，而不是只改画布", 2)
    add_body(doc, "我没有把人工修改做成单纯拖拽 SVG 图形，而是让元素和关系的结构化数据发生变化。这样 SysML textual view、模型画布、预评审和入库结果都能保持一致。")

    add_heading(doc, "5.5 预评审不能只靠大模型主观判断", 2)
    add_body(doc, "校验功能如果完全依赖大模型，会出现不稳定和不可复现的问题。因此我先实现确定性规则，再把大模型定位为后续的解释、补充建议和复杂修复助手。")

    add_heading(doc, "6. 当前系统的最终闭环", 1)
    add_callout(
        doc,
        "最终闭环",
        "我现在搭建出的系统闭环是：项目工作区承载上下文，RAG 提供知识支撑，大模型生成候选 SysML，设计师进行人工修改，系统执行预评审并一键修复，设计师采纳后模型入库，再推送 MagicDraw 或进入仿真。这个闭环把 AI 的效率、人为确认的可靠性和系统工程工具链连接了起来。",
    )

    add_heading(doc, "7. 后续我会优先扩展的方向", 1)
    for item in [
        "把当前前端权限控制升级为后端真实用户、角色、项目授权和操作审计。",
        "把轻量 RAG 检索升级为向量数据库，并增加召回率测试集和 rerank。",
        "增加模型库管理页面，直接查看已入库模型、版本、关系和采纳记录。",
        "把预评审规则配置化，允许不同项目使用不同建模规范。",
        "把 model_feedback_logs 导出为训练数据，用于后续大模型微调。",
        "增强 MagicDraw 双向同步，让 MagicDraw 中的人工修改也能回写到系统草稿或模型库。",
        "让仿真结果反向更新模型约束和验证报告，形成设计-建模-仿真-验证闭环。",
    ]:
        add_bullet(doc, item)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_doc())
