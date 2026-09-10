from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "AI_MBSE_系统整体框架与流程说明.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
GRAY = RGBColor(89, 101, 121)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
BORDER = "D8E2EF"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_font(paragraph, size: float = 11, color: RGBColor | None = None, bold: bool | None = None):
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None, size: float = 9.4):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def set_table_geometry(table, widths_in: list[float]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            if idx >= len(row.cells):
                continue
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(1, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)


def add_rule(paragraph, color="D8E2EF"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_heading(doc, text: str, level: int):
    p = doc.add_heading(text, level=level)
    if level == 1:
        set_paragraph_font(p, 16, BLUE, True)
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_paragraph_font(p, 13, BLUE, True)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        set_paragraph_font(p, 12, DARK_BLUE, True)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, 11, True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, 11)
    else:
        r = p.add_run(text)
        set_run_font(r, 11)
    return p


def add_bullet(doc, text: str, level: int = 0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, 10.5)
    return p


def add_number(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, 10.5)
    return p


def add_callout(doc, title: str, body: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(title)
    set_run_font(r1, 10.5, True, NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(body)
    set_run_font(r2, 10.2, False, None)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float], header_fill: str = LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        shade_cell(table.cell(0, i), header_fill)
        set_cell_text(table.cell(0, i), header, bold=True, color=NAVY, size=9.2)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, size=8.9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name in ["List Bullet", "List Bullet 2", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("AI-MBSE 系统功能框架说明")
    set_run_font(run, 9, False, GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("AI-MBSE-Demo 原型文档")
    set_run_font(r, 9, False, GRAY)
    return doc


def add_title_block(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("AI 赋能 MBSE 系统功能框架与流程说明")
    set_run_font(r, 23, True, NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("基于当前 AI-MBSE-Demo 原型的总体架构、业务流程与模块搭建方法")
    set_run_font(r, 13.5, False, GRAY)

    metadata = [
        ("文档类型", "技术架构说明 / 功能建设总结"),
        ("适用场景", "课题汇报、系统方案、论文实现章节、后续开发交接"),
        ("系统范围", "人员权限、RAG 知识库、大模型交互、SysML 生成、人为回路、预评审校验、模型入库、MagicDraw 同步"),
        ("版本日期", date.today().isoformat()),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1.35, 5.15])
    for idx, (label, value) in enumerate(metadata):
        shade_cell(table.cell(idx, 0), LIGHT_GRAY)
        set_cell_text(table.cell(idx, 0), label, bold=True, color=NAVY)
        set_cell_text(table.cell(idx, 1), value)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    add_rule(rule)


def build_doc():
    doc = setup_document()
    add_title_block(doc)

    add_heading(doc, "1. 文档摘要", 1)
    add_body(
        doc,
        "本系统是一个面向 MBSE 建模过程的 AI 辅助设计平台。它把设计师输入、知识库检索、大模型问答、SysML 模型生成、人为确认、模型预评审和模型入库组织成一条闭环流程。",
    )
    add_body(
        doc,
        "核心原则：AI 生成的模型不直接进入正式模型库，而是先形成候选草稿；设计师可以人工修改，系统可以执行预评审并一键修复；只有在用户点击“采纳”后，当前草稿才写入模型库并可继续推送 MagicDraw。",
        bold_prefix="核心原则：",
    )
    add_callout(
        doc,
        "当前原型定位",
        "系统已经具备人员/项目侧栏、RAG 知识库构建、大模型流式对话、SysML 候选模型生成、人工修改、采纳/拒绝反馈、预评审校验、一键修复、模型入库、MagicDraw 同步与仿真入口等能力。",
    )

    add_heading(doc, "2. 总体框架", 1)
    add_body(doc, "整体架构可以理解为七层：前端交互层、应用编排层、AI 与知识增强层、草稿与人为回路层、建模服务层、数据持久化层、外部工具集成层。")
    add_table(
        doc,
        ["层级", "主要组成", "职责"],
        [
            ["前端交互层", "项目栏、AI 智能体、SysML textual view、模型画布、预评审弹窗", "承载设计师的输入、查看、修改、校验、采纳和推送操作。"],
            ["应用编排层", "FastAPI routers、前端 API client、工作区状态", "把用户操作路由到聊天、文档、SysML、MagicDraw、仿真等服务。"],
            ["AI 与知识增强层", "LLMService、RAGService、知识库文档切片、轻量检索", "根据对话和知识库材料回答问题，并为建模生成提供上下文。"],
            ["草稿与人为回路层", "model_drafts、model_feedback_logs、人工修改组件", "保存 AI 初稿与人工修订版，记录采纳、拒绝、修正、修复等反馈。"],
            ["建模服务层", "ModelGenerationService、SysMLAdapter、DraftValidationService", "生成 SysML 元素/关系/文本，执行预评审和一键修复，完成采纳入库。"],
            ["数据持久化层", "SQLite 数据库 ai_mbse_demo.db", "保存知识片段、草稿、反馈、正式模型元素、关系和校验历史。"],
            ["外部工具集成层", "SysML v2 API、MagicDraw Bridge、仿真接口", "与外部建模平台、MagicDraw 和仿真工具进行数据交换。"],
        ],
        [1.15, 2.1, 3.25],
        LIGHT_BLUE,
    )

    add_heading(doc, "3. 端到端业务流程", 1)
    workflow_steps = [
        "设计师进入系统，选择或新建项目。不同角色可拥有不同项目范围和操作权限。",
        "设计师在 AI 智能体输入问题、需求或方案描述；可上传临时附件，也可上传知识库文档。",
        "系统根据模式选择直接问答或知识库增强问答；RAG 服务检索文档片段、模型图谱和内置知识。",
        "设计师点击生成，系统把对话上下文和需求输入转化为 SysML 候选工程，生成元素、关系和 SysML textual view。",
        "候选模型进入预览与确认界面，不直接入库；设计师可以编辑元素、修改/新增/删除连接线。",
        "设计师点击执行模型校验，系统生成《预评审报告》，按错误、警告、提示列出问题和修复建议。",
        "对可自动修复的问题，设计师可点击一键修复；系统更新草稿、重算 SysML textual view，并记录修复反馈。",
        "设计师确认无误后点击采纳；系统把当前草稿写入模型库，并记录正反馈。",
        "采纳后的模型可以推送 MagicDraw、导出交换包或进入仿真流程。",
    ]
    for step in workflow_steps:
        add_number(doc, step)

    add_heading(doc, "4. 关键数据对象与流转", 1)
    add_table(
        doc,
        ["对象", "产生位置", "主要用途", "是否正式入库"],
        [
            ["临时附件", "对话输入区纸夹上传", "只用于当前轮上下文，不进入知识库。", "否"],
            ["知识库文档", "知识库“上传入库”", "解析、切片并参与 RAG 检索。", "进入 documents / knowledge_chunks"],
            ["AI 候选模型", "SysML 生成接口", "作为预览草稿供设计师修改。", "否，先进入 model_drafts"],
            ["当前草稿", "人工修改或一键修复后", "作为采纳时的真实模型来源。", "采纳后进入正式模型库"],
            ["反馈日志", "采纳、拒绝、修正、修复操作", "用于追踪人为回路和后续微调数据。", "进入 model_feedback_logs"],
            ["正式模型", "点击采纳后", "供模型库、图谱、MagicDraw、仿真使用。", "进入 sysml_elements / sysml_relationships"],
        ],
        [1.25, 1.55, 2.65, 1.05],
    )

    add_heading(doc, "5. 功能模块说明", 1)

    add_heading(doc, "5.1 人员权限管理与项目工作区", 2)
    add_body(doc, "该模块负责把用户操作限定在对应工作区内，并把对话、草稿、同步结果、校验报告与项目绑定。当前原型已具备项目侧栏、新建项目、项目重命名、受保护基线项目等能力。")
    for text in [
        "角色与工作区：区分设计师、管理员等角色；设计师通常只能操作授权项目。",
        "项目级状态：每个项目保存独立的聊天记录、草稿模型、MagicDraw 同步状态、预评审报告。",
        "权限表现：受保护项目不能被普通设计师删除；后续可扩展为后端鉴权和审计。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "搭建步骤", 3)
    for text in [
        "定义角色、项目、工作区三个基础概念。",
        "在前端建立项目侧栏和项目级状态容器。",
        "把聊天记录、模型草稿、预评审报告都挂到当前项目下。",
        "对删除、改名、采纳、推送等高风险操作增加角色判断。",
        "后续接入真实用户表、Token、角色权限矩阵和操作审计表。",
    ]:
        add_number(doc, text)

    add_heading(doc, "5.2 RAG 知识库构建", 2)
    add_body(doc, "RAG 模块用于把参考文献、规范文档、设计资料转化为可检索知识。系统区分“临时附件”和“知识库文档”：临时附件只参与当前对话，知识库文档会持久化并长期参与检索。")
    add_table(
        doc,
        ["子功能", "当前实现", "后续增强"],
        [
            ["文档上传", "支持知识库上传和临时附件上传两条入口。", "增加文档分类、版本、权限范围。"],
            ["PDF 解析", "使用 pypdf 提取文本，避免 PDF 只产生 1 个片段。", "加入 OCR、版面结构识别、表格抽取。"],
            ["知识切片", "把文档分块写入 knowledge_chunks。", "引入语义切片和章节层级索引。"],
            ["检索召回", "轻量检索文档片段、模型图谱和种子知识。", "升级为 FAISS/Milvus/Qdrant + rerank。"],
            ["知识库管理", "显示文档数/片段数，支持手动重置。", "增加文档级删除、召回率评测面板。"],
        ],
        [1.25, 2.75, 2.5],
    )
    add_heading(doc, "搭建步骤", 3)
    for text in [
        "建立 documents 和 knowledge_chunks 表，保存原文和切片。",
        "实现文档上传接口，按文件类型抽取文本。",
        "实现切片策略，记录 chunk_index、title、keywords。",
        "在聊天服务中加入 RAG 检索，把高相关片段拼入大模型上下文。",
        "前端增加知识库状态条、上传入库、刷新和重置按钮。",
        "通过标准问题集测试召回率，逐步优化切片、关键词和向量检索策略。",
    ]:
        add_number(doc, text)

    add_heading(doc, "5.3 大模型交互与智能体", 2)
    add_body(doc, "AI 智能体承担需求澄清、知识问答、建模意图识别和生成提示构造。系统提供直接问答、知识库问答和自动识别模式，并支持流式输出。")
    for text in [
        "直接问答：只调用大模型，不强制使用知识库。",
        "知识库问答：检索知识片段后再回答，并展示证据来源。",
        "自动识别：根据用户问题判断是否需要知识库、是否触发建模生成。",
        "流式输出：前端逐步显示回答，提高交互体验。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "搭建步骤", 3)
    for text in [
        "封装 LLMService，统一模型 API 地址、Key、模型名称和异常处理。",
        "实现 chat 和 chat/stream 接口，支持普通与流式响应。",
        "在 prompt 中加入角色、阶段、工作流约束、附件摘要和 RAG 上下文。",
        "实现建模意图分类，避免普通问答误触发 SysML 生成。",
        "在前端保存对话历史，并把最近上下文传给生成接口。",
    ]:
        add_number(doc, text)

    add_heading(doc, "5.4 SysML 模型生成与预览", 2)
    add_body(doc, "模型生成模块把对话和需求转化为结构化模型：elements、relationships、sysml_text。生成结果首先保存为草稿，展示在 SysML textual view 和模型画布中。")
    add_table(
        doc,
        ["输出", "含义", "示例"],
        [
            ["elements", "需求、模块、接口、活动、约束等模型元素。", "REQ-001、BLK-001、IF-001"],
            ["relationships", "元素之间的追溯、满足、分配、验证等关系。", "satisfy、trace、allocate、verify"],
            ["sysml_text", "根据结构化模型渲染出的 SysML textual view。", "package、requirement def、part def"],
            ["draft_id", "后台草稿编号，用于后续修改、校验、采纳。", "草稿 #6"],
        ],
        [1.3, 3.1, 2.1],
    )
    add_heading(doc, "搭建步骤", 3)
    for text in [
        "定义 ModelElement 和 ModelRelationship 结构。",
        "实现 ModelGenerationService：优先解析 LLM 结构化输出，失败时使用动态或行业模板兜底。",
        "把生成结果渲染成 SysML textual view。",
        "生成后创建 model_drafts 记录，保存 original_json 和 current_json。",
        "前端把 textual view 和画布绑定到同一份 projectResult，保证联动显示。",
    ]:
        add_number(doc, text)

    add_heading(doc, "5.5 人为回路确认与反馈记录", 2)
    add_body(doc, "人为回路模块确保 AI 不直接决定模型入库。设计师可以修改元素、修改关系、新增/删除连接线、拒绝候选方案或采纳当前草稿。所有操作都写入反馈日志。")
    for text in [
        "元素编辑：修改名称、类型、包路径、来源需求、状态和描述。",
        "关系编辑：修改起点、终点、关系类型和说明。",
        "连接线增删：新增关系或删除无效关系。",
        "采纳/拒绝：采纳会入库，拒绝会记录负反馈。",
        "反馈日志：保存 before_json 和 after_json，可用于后续大模型微调。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "搭建步骤", 3)
    for text in [
        "新增 model_drafts 表，区分 AI 原始结果 original_json 和当前草稿 current_json。",
        "新增 model_feedback_logs 表，记录 revise、add、delete、accept、reject、fix 等动作。",
        "实现草稿元素和关系的增删改接口。",
        "前端提供人工修改面板，保存后替换当前草稿并重算 SysML 文本。",
        "采纳时使用 current_json 入库，确保入库的是人工确认后的模型。",
    ]:
        add_number(doc, text)

    add_heading(doc, "5.6 模型预评审与智能化校验", 2)
    add_body(doc, "预评审模块面向候选草稿执行规则检查，生成《预评审报告》。系统先使用确定性规则保证稳定性，再为可修复问题提供一键修复按钮。")
    add_table(
        doc,
        ["规则类别", "检查内容", "自动修复策略"],
        [
            ["结构一致性", "关系端点是否存在、ID 是否重复。", "删除无效关系或提示人工确认。"],
            ["SysML 规范", "关系类型是否合法、包路径是否匹配类型。", "规范化关系类型和包路径。"],
            ["建模规范", "需求是否被 satisfy，接口是否 allocate，约束是否 verify。", "自动补充建议关系。"],
            ["属性规范", "功率、带宽、频率、时延等属性是否有类型和单位。", "补充 value_type 和 unit。"],
            ["发布状态", "元素是否仍为 candidate。", "规范化为 reviewed。"],
            ["文本同步", "SysML textual view 是否与结构化模型一致。", "重新渲染 SysML 文本。"],
        ],
        [1.25, 3.0, 2.25],
        LIGHT_BLUE,
    )
    add_heading(doc, "搭建步骤", 3)
    for text in [
        "新增 DraftValidationService，读取 model_drafts.current_json。",
        "定义 DraftValidationIssue 和 DraftValidationReport 结构。",
        "实现 validate 接口，返回评分、结论、统计和问题列表。",
        "为可修复问题绑定 fix_action，例如 normalize_package、add_relationship、set_attribute_type。",
        "实现 fix 接口，按 issue_ids 修改草稿、重算 SysML 文本并记录 action=fix 的反馈日志。",
        "前端增加“执行模型校验”按钮和《预评审报告》弹窗，支持单项修复和全部修复。",
    ]:
        add_number(doc, text)

    add_heading(doc, "5.7 模型入库与模型库", 2)
    add_body(doc, "模型库当前对应后端 SQLite 数据库中的正式 SysML 模型表。只有点击采纳后，当前草稿才会通过 SysMLAdapter 写入 sysml_elements 和 sysml_relationships。")
    add_table(
        doc,
        ["表/存储", "内容", "说明"],
        [
            ["model_drafts", "草稿模型 original_json / current_json", "候选阶段使用，不代表正式入库。"],
            ["model_feedback_logs", "采纳、拒绝、修改、修复记录", "用于审计和微调数据积累。"],
            ["sysml_elements", "正式模型元素", "采纳后写入。"],
            ["sysml_relationships", "正式模型关系", "采纳后写入。"],
            ["validation_reports", "历史校验报告", "已有入库模型校验使用；草稿预评审使用新接口即时返回。"],
            ["外部 SysML API", "远程项目和提交", "配置 SYSML_ENABLE_REAL_API=true 时尝试调用。"],
        ],
        [1.65, 2.35, 2.5],
    )
    add_heading(doc, "搭建步骤", 3)
    for text in [
        "实现 SysMLAdapter.commit_model，把元素和关系 upsert 到正式表。",
        "配置真实 SysML API 时，优先向远程 SysML v2 项目创建提交。",
        "真实 API 不可用时，降级到本地 SQLite mock 保存。",
        "前端限制未采纳草稿不能推送 MagicDraw，避免绕过确认流程。",
        "后续可增加模型库管理页面，展示已入库元素、关系、版本和采纳记录。",
    ]:
        add_number(doc, text)

    add_heading(doc, "5.8 MagicDraw 同步、导出与仿真", 2)
    add_body(doc, "MagicDraw 同步位于采纳之后，负责把正式模型或交换包推送到 MagicDraw Bridge。仿真模块当前提供统一入口，可扩展到 GMAT、Simulink 等工具。")
    for text in [
        "MagicDraw 状态条显示 Bridge 是否连接。",
        "推送动作读取已采纳模型，生成 exchange_payload 或调用 Bridge。",
        "导出功能把当前工程包保存为 JSON，用于离线交付。",
        "仿真入口可基于模型 ID 调用后端仿真服务。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "6. 主要接口与服务映射", 1)
    add_table(
        doc,
        ["能力", "接口/服务", "说明"],
        [
            ["知识库上传", "POST /api/documents/upload", "上传文档、解析文本、切片入库。"],
            ["知识库状态", "GET /api/documents/knowledge-summary", "返回文档数、片段数和最新文档。"],
            ["知识库重置", "DELETE /api/documents/knowledge", "清空知识库文档和索引片段。"],
            ["流式对话", "POST /api/chat/stream", "大模型流式问答，附带工具调用和知识证据。"],
            ["生成草稿", "POST /api/sysml/generate-project", "生成 SysML 候选工程并创建 draft。"],
            ["修改元素", "PATCH /api/sysml/drafts/{id}/elements/{element_id}", "更新草稿元素并记录反馈。"],
            ["修改关系", "PATCH /api/sysml/drafts/{id}/relationships/{rel_id}", "更新草稿连接线。"],
            ["新增关系", "POST /api/sysml/drafts/{id}/relationships", "新增草稿连接线。"],
            ["删除关系", "DELETE /api/sysml/drafts/{id}/relationships/{rel_id}", "删除草稿连接线。"],
            ["草稿校验", "POST /api/sysml/drafts/{id}/validate", "生成预评审报告。"],
            ["一键修复", "POST /api/sysml/drafts/{id}/fix", "按 issue 修复草稿并重算 SysML。"],
            ["采纳入库", "POST /api/sysml/drafts/{id}/accept", "当前草稿写入正式模型库。"],
            ["推送 MagicDraw", "POST /api/magicdraw/publish", "生成交换包或调用 Bridge。"],
            ["运行仿真", "POST /api/simulation/run", "执行后端仿真任务。"],
        ],
        [1.25, 2.65, 2.6],
        LIGHT_GRAY,
    )

    add_heading(doc, "7. 分阶段搭建路线", 1)
    add_body(doc, "如果从零复现该系统，可以按以下阶段推进，每一阶段都形成可演示闭环，再继续叠加下一层能力。")
    phases = [
        ("阶段 1：基础框架", "搭建 React/Vite 前端、FastAPI 后端、SQLite 数据库、健康检查接口和统一 API client。"),
        ("阶段 2：项目与权限", "实现项目列表、工作区隔离、角色判断、受保护基线项目和项目级状态。"),
        ("阶段 3：大模型对话", "接入 LLMService、普通聊天、流式输出、模式选择和错误提示。"),
        ("阶段 4：知识库 RAG", "实现知识库上传、PDF 解析、切片入库、检索服务、知识库状态和重置。"),
        ("阶段 5：SysML 生成", "定义模型元素/关系结构，生成候选草稿，渲染 textual view 和画布。"),
        ("阶段 6：人为回路", "增加草稿表、反馈表、人工编辑、采纳/拒绝和采纳后入库。"),
        ("阶段 7：预评审校验", "实现草稿校验规则、报告弹窗、可修复 issue 和一键修复。"),
        ("阶段 8：工具集成", "接入 SysML API、MagicDraw Bridge、导出交换包和仿真接口。"),
        ("阶段 9：评测与治理", "加入 RAG 召回率评测、校验规则覆盖率、反馈数据导出和权限审计。"),
    ]
    for title, body in phases:
        add_number(doc, f"{title}：{body}")

    add_heading(doc, "8. 当前实现边界与后续建议", 1)
    add_body(doc, "当前系统已经形成完整原型闭环，但仍有若干工程化增强方向。")
    add_table(
        doc,
        ["方向", "建议"],
        [
            ["权限管理", "从前端角色判断升级为后端用户、角色、项目授权和操作审计。"],
            ["RAG 数据库", "把 SQLite 轻量检索升级为向量数据库，并加入 rerank 和召回率评测集。"],
            ["知识图谱", "将已采纳模型和知识文档实体统一进入图数据库，支持影响分析和关系推理。"],
            ["模型库页面", "新增模型库管理页面，查看已采纳模型、版本、关系图、反馈历史。"],
            ["预评审规则", "把规则配置化，支持项目级建模规范、严重级别、自动修复策略开关。"],
            ["微调数据", "把 model_feedback_logs 转换为训练样本，用于模型生成和修复策略优化。"],
            ["版本管理", "对草稿和正式模型增加版本号、差异对比、回滚和审批链。"],
            ["外部工具", "增强 MagicDraw 双向同步，并把仿真结果反写到模型约束和验证报告中。"],
        ],
        [1.35, 5.15],
    )

    add_heading(doc, "9. 一句话总结", 1)
    add_callout(
        doc,
        "系统闭环",
        "AI-MBSE-Demo 的核心不是“让大模型直接画图”，而是构建一个可控的系统工程闭环：知识增强理解需求，AI 生成候选模型，设计师人工确认，系统自动预评审和修复，最终把确认后的模型入库并同步到专业工具。",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build_doc()
    print(path)
