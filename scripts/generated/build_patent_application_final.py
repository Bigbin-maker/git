from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\ai_mbse1\AI-MBSE-Demo")
OUT = ROOT / "docs" / "generated" / "一种面向复杂系统工程模型的知识增强式生成校验与变更治理方法及系统_发明专利申请文件初稿.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(80, 80, 80)
BORDER = "D9DEE7"
FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"


def set_run_font(run, size=11, bold=False, color=None, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def set_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_para(doc, text="", size=11, bold=False, color=None, align=None, before=0, after=6, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12 if level == 2 else 8)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6 if level == 2 else 4)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, fill=FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        shade_cell(cell, fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i == 0 and len(text) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_run_font(r, size=10)
    add_para(doc, "", after=4)


def set_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    r = header.add_run("发明专利申请文件初稿")
    set_run_font(r, size=9, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("内部草案 - 提交前请由专利代理人复核")
    set_run_font(r, size=9, color=GRAY)


def add_claims(doc):
    add_heading(doc, "权利要求书", 1)
    claims = [
        "1. 一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法，其特征在于，包括如下步骤：获取复杂系统设计过程中的多源工程数据；从所述多源工程数据中抽取工程对象，所述工程对象至少包括需求、功能、模块、接口、参数、约束和验证项；基于所述工程对象构建工程知识底座，所述工程知识底座包括向量语义索引、工程知识图谱和工程规则库；响应于用户输入的建模任务，基于所述工程知识底座执行知识增强检索，形成与所述建模任务对应的生成上下文和来源证据；基于所述生成上下文生成候选工程模型，所述候选工程模型包括模型元素、模型关系和证据链；对所述候选工程模型执行模型校验，得到校验结果、问题定位和修复建议；接收具备权限的用户对所述候选工程模型的审核操作，并根据所述审核操作更新候选工程模型的状态；响应于工程模型变更请求，在沙箱环境中执行变更影响分析，得到影响范围、冲突信息和处置建议；在所述候选工程模型或变更结果满足预设入库条件后，将其合并至发布分支并生成审计记录。",
        "2. 根据权利要求1所述的方法，其特征在于，所述多源工程数据包括需求文档、接口控制文件、历史工程模型、工程建模工具导出文件、仿真结果、评审记录、版本记录、权限配置数据和用户反馈数据中的一种或多种。",
        "3. 根据权利要求1所述的方法，其特征在于，从所述多源工程数据中抽取工程对象包括：识别自然语言文本中的需求描述、功能动作、系统组成、外部参与者、接口名称、接口方向、参数名称、参数取值、约束条件、验证方式和生命周期状态，并为每一工程对象建立唯一标识、对象类型、来源位置和可信度标记。",
        "4. 根据权利要求1所述的方法，其特征在于，所述工程知识图谱包括需求与模型元素之间的满足关系、需求之间的派生关系、模块之间的组成关系、接口之间的连接关系、参数与约束之间的约束关系、验证项与需求之间的验证关系、模型元素之间的依赖关系以及工程模型不同版本之间的变更关系。",
        "5. 根据权利要求1所述的方法，其特征在于，基于所述工程知识底座执行知识增强检索包括：基于所述向量语义索引召回与所述建模任务语义相关的文档片段；基于所述工程知识图谱召回与目标工程对象相邻的模型元素和关系；基于所述工程规则库召回与所述建模任务相关的建模规则、命名规则、接口规则或校验规则；对召回结果进行去重、排序和证据标注，形成生成上下文。",
        "6. 根据权利要求1所述的方法，其特征在于，所述候选工程模型包括文本化工程模型、图形化模型视图、模型元素清单、模型关系清单、接口连接清单、约束清单以及所述候选工程模型与来源证据之间的追溯关系。",
        "7. 根据权利要求1所述的方法，其特征在于，所述工程模型包括SysML模型、UML模型、AADL模型、Simulink接口模型、数字孪生模型或企业自定义工程模型中的一种或多种。",
        "8. 根据权利要求1所述的方法，其特征在于，对所述候选工程模型执行模型校验包括：模型语言语法校验、模型语义一致性校验、接口完整性校验、命名规范校验、属性类型校验、连接合法性校验、约束一致性校验、需求覆盖性校验和仿真可满足性校验中的一种或多种。",
        "9. 根据权利要求8所述的方法，其特征在于，当所述校验结果表明候选工程模型存在可自动修复问题时，生成与所述问题定位对应的模型修复补丁；在接收到用户确认操作后，将所述模型修复补丁应用于所述候选工程模型，并重新执行模型校验。",
        "10. 根据权利要求1所述的方法，其特征在于，所述审核操作包括采纳、局部修正、驳回、提交复核、备注和权限审批；所述审核操作被记录为反馈数据，所述反馈数据包括用户身份、操作时间、被操作模型元素、修改前后差异、审核意见和反馈类型。",
        "11. 根据权利要求1所述的方法，其特征在于，在沙箱环境中执行变更影响分析包括：在不改变发布分支的条件下模拟模型元素、接口、参数或约束的变更；基于工程知识图谱和工程规则库进行影响传播；识别直接影响对象、间接影响对象、受影响验证项、潜在冲突项和建议处理措施；输出影响拓扑图、影响矩阵或冲突清单。",
        "12. 根据权利要求1所述的方法，其特征在于，所述候选工程模型在入库前具有候选状态，发布分支中的工程模型元素具有权威状态；当候选工程模型经审核通过并完成合并后，将其状态由候选状态转换为权威状态。",
        "13. 根据权利要求1所述的方法，其特征在于，还包括：调用外部工程建模工具进行模型同步，和/或调用外部仿真工具执行与目标需求相关的仿真任务；接收所述外部工程建模工具或外部仿真工具返回的模型同步结果或仿真结果；将所述模型同步结果或仿真结果作为证据关联至对应的需求、模型元素、参数、约束或验证项。",
        "14. 一种面向复杂系统工程模型的知识增强式生成、校验与变更治理系统，其特征在于，包括：数据接入模块，用于获取复杂系统设计过程中的多源工程数据；工程对象抽取模块，用于从所述多源工程数据中抽取工程对象；知识底座模块，用于构建并维护包括向量语义索引、工程知识图谱和工程规则库的工程知识底座；知识增强检索模块，用于根据建模任务检索生成上下文和来源证据；模型生成模块，用于生成候选工程模型；模型校验模块，用于对候选工程模型执行校验并生成校验结果和修复建议；人在回路模块，用于接收用户对候选工程模型的审核操作；变更治理模块，用于在沙箱环境中执行变更影响分析；分支合并模块，用于将通过审核的候选工程模型或变更结果合并至发布分支；审计反馈模块，用于记录审计信息并将审核反馈回写至工程知识底座。",
        "15. 根据权利要求14所述的系统，其特征在于，还包括工具联动模块，所述工具联动模块用于与工程建模工具、仿真工具、图数据库、向量数据库或权限认证系统中的一种或多种进行数据交互。",
        "16. 一种电子设备，包括处理器和存储器，所述存储器中存储有计算机程序，其特征在于，所述计算机程序被所述处理器执行时实现权利要求1至13任一项所述的方法。",
        "17. 一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至13任一项所述的方法。",
    ]
    for claim in claims:
        add_para(doc, claim, after=8)


def add_description(doc):
    add_heading(doc, "说明书", 1)
    add_heading(doc, "一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统", 2)

    add_heading(doc, "技术领域", 2)
    add_para(doc, "本发明涉及复杂系统工程、模型驱动工程、人工智能辅助设计、工程知识管理和模型变更治理技术领域，尤其涉及一种利用知识增强机制实现工程模型生成、模型校验、人在回路审核、变更影响分析和分支合并入库的方法及系统。")

    add_heading(doc, "背景技术", 2)
    for text in [
        "复杂系统工程设计通常涉及需求、功能、逻辑架构、物理架构、接口、参数、约束、验证项、仿真结果和评审意见等多类工程数据。上述数据通常分散在需求文档、建模工具、图数据库、接口控制文档、仿真工具、评审记录和项目管理系统中，导致工程人员在模型构建、模型审查和模型变更时需要进行大量人工检索、比对和确认。",
        "随着大语言模型技术的发展，利用人工智能辅助工程建模成为可能。然而，通用大语言模型直接生成工程模型时，容易出现与项目上下文不一致、来源不可追溯、接口关系缺失、属性类型不完整、约束不满足以及无法满足工程入库要求等问题。单纯的文本问答或文档检索不能直接解决工程模型从生成到校验、审核、变更和入库的治理问题。",
        "现有模型驱动工程工具通常能够提供模型编辑、模型可视化和部分规则检查能力，但其对多源工程知识融合、知识增强模型生成、候选模型状态管理、人工审核反馈沉淀、沙箱式变更预演以及分支合并入库的支持不足。特别是在多人协同、模型持续演进和工程数据需要权威发布的场景下，亟需一种能够将知识增强生成、自动校验、人在回路审核、变更影响分析和数据入库管理贯通的通用技术方案。",
    ]:
        add_para(doc, text)

    add_heading(doc, "发明内容", 2)
    add_heading(doc, "要解决的技术问题", 3)
    for text in [
        "解决通用人工智能生成工程模型时缺少项目知识约束和工程来源证据的问题。",
        "解决候选工程模型入库前缺乏自动校验、修复建议和人工审核闭环的问题。",
        "解决工程模型变更对需求、接口、参数、约束、验证项和仿真结果的影响难以提前识别的问题。",
        "解决AI生成模型、人工修正结果和发布分支权威模型之间缺乏状态治理和审计追踪的问题。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "技术方案", 3)
    add_para(doc, "为解决上述技术问题，本发明提供一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法。该方法首先接入多源工程数据，并从中抽取需求、功能、模块、接口、参数、约束和验证项等工程对象；随后构建包括向量语义索引、工程知识图谱和工程规则库的工程知识底座；当用户发起建模任务时，系统基于工程知识底座进行知识增强检索，形成生成上下文和来源证据，并生成候选工程模型；候选工程模型经过模型校验、修复建议生成、人在回路审核后，才能进入分支合并和发布入库流程；当模型发生变更时，系统在沙箱环境中执行变更影响分析，以识别影响范围和冲突信息。")
    add_para(doc, "本发明还提供一种系统，包括数据接入模块、工程对象抽取模块、知识底座模块、知识增强检索模块、模型生成模块、模型校验模块、人在回路模块、变更治理模块、分支合并模块、审计反馈模块以及工具联动模块。")

    add_heading(doc, "有益效果", 3)
    for text in [
        "本发明通过向量语义索引、工程知识图谱和工程规则库共同约束生成过程，使候选工程模型能够绑定项目上下文和来源证据，降低模型生成过程中的不一致和不可追溯风险。",
        "本发明将候选模型生成、自动校验、修复建议和人工审核统一在入库前流程中，使AI生成内容不会直接污染发布分支，提升工程模型质量。",
        "本发明通过沙箱式变更影响分析，在不改变权威发布模型的前提下模拟参数、接口或模型元素变更，提前识别直接影响、间接影响和潜在冲突。",
        "本发明通过候选状态、权威状态、分支合并和审计记录机制，实现AI生成内容从候选模型到权威模型的可控转化。",
        "本发明具有跨模型语言扩展能力，可适用于SysML、UML、AADL、Simulink接口模型、数字孪生模型和企业自定义工程模型。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "附图说明", 2)
    for text in [
        "图1为本发明方法的总体流程图。",
        "图2为本发明系统的功能模块架构图。",
        "图3为工程知识底座的数据结构示意图。",
        "图4为知识增强候选工程模型生成流程图。",
        "图5为候选工程模型校验与人在回路审核流程图。",
        "图6为沙箱式变更影响分析与分支合并入库流程图。",
        "图7为需求、工程模型、仿真证据和审计记录之间的追溯关系示意图。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "具体实施方式", 2)
    add_heading(doc, "实施例一：总体流程", 3)
    add_para(doc, "在一个实施例中，本发明应用于低轨通信卫星、无人系统、航空装备、工业生产线或其他复杂系统的工程设计过程。系统接收用户上传的需求说明、接口控制文件、已有工程模型和仿真结果文件，将其解析为统一工程对象，并建立需求、功能、模块、接口、参数、约束、验证项之间的关系。用户发起建模任务时，系统先检索相关文档片段、图谱邻域和规则条目，再生成候选工程模型；候选工程模型经过自动校验和人工审核后，才可进入开发分支或发布分支。")
    add_table(
        doc,
        ["步骤", "处理内容", "输出结果"],
        [
            ("S1", "接入需求文档、接口规范、模型文件、仿真结果、评审记录和版本数据。", "多源工程数据集合。"),
            ("S2", "抽取需求、功能、模块、接口、参数、约束和验证项等工程对象。", "结构化工程对象。"),
            ("S3", "构建向量语义索引、工程知识图谱和工程规则库。", "工程知识底座。"),
            ("S4", "根据建模任务进行知识增强检索。", "生成上下文和来源证据。"),
            ("S5", "生成候选工程模型。", "模型元素、模型关系和证据链。"),
            ("S6", "执行模型校验和修复建议生成。", "校验结果、问题定位和修复补丁。"),
            ("S7", "由设计师或审核人进行采纳、修正、驳回或复核。", "审核记录和候选模型状态。"),
            ("S8", "在沙箱环境中执行变更影响分析。", "影响拓扑、影响矩阵和冲突清单。"),
            ("S9", "通过分支合并入库并生成审计记录。", "发布分支权威工程模型。"),
        ],
        [900, 5200, 3260],
        BLUE_FILL,
    )

    add_heading(doc, "实施例二：知识增强生成", 3)
    add_para(doc, "用户输入建模任务后，系统根据当前项目、用户角色、当前分支、上传附件、历史对话和已有模型状态形成任务上下文。知识增强检索模块分别从向量语义索引中召回相似文档片段，从工程知识图谱中召回相关模型元素及其邻域关系，从工程规则库中召回建模规则和校验规则。上述召回结果经过去重、排序和证据标注后，形成生成上下文包。")
    add_para(doc, "模型生成模块基于生成上下文包输出候选工程模型。候选工程模型可以包括文本化模型描述、图形化模型视图、模型元素清单、接口连接清单、约束清单以及来源证据链。对于SysML实施方式，候选工程模型可以表现为需求图、用例图、内部模块图、活动图或接口关系图；对于其他模型语言，也可以转换为UML、AADL、Simulink接口模型或企业自定义模型。")

    add_heading(doc, "实施例三：模型校验与修复", 3)
    add_para(doc, "模型校验模块对候选工程模型执行多层校验。第一层为语法校验，用于检查模型语言语法、元素定义和关系定义是否合法；第二层为工程语义校验，用于检查需求与模块、接口与连接、约束与参数、验证项与需求之间是否一致；第三层为项目规范校验，用于检查命名、编号、属性类型、生命周期状态和权限边界是否符合项目规则；第四层为仿真可满足性校验，用于判断关键需求是否关联了仿真场景、仿真参数和仿真证据。")
    add_para(doc, "当校验发现命名不规范、属性类型缺失、连接方向错误、接口引用缺失等可自动修复问题时，系统生成模型修复补丁，并在用户确认后执行规范化修复；当问题涉及工程判断时，系统仅给出建议，并交由具备权限的人员在审核环节处理。")

    add_heading(doc, "实施例四：人在回路审核", 3)
    add_para(doc, "候选工程模型默认处于候选状态，不能直接成为发布分支中的权威模型。具有相应权限的设计师或审核人可以对候选模型执行采纳、局部修正、驳回、提交复核或备注操作。系统记录用户身份、操作时间、被操作模型元素、修改前后差异、审核意见和反馈类型。上述反馈数据可回写至工程知识底座，用于后续检索排序、规则优化、提示词优化或模型微调。")

    add_heading(doc, "实施例五：沙箱变更影响分析与合并入库", 3)
    add_para(doc, "当用户修改模型元素、接口、参数或约束时，系统不直接覆盖发布分支，而是在开发分支或沙箱环境中生成变更预演。变更治理模块基于工程知识图谱中的依赖关系、满足关系、连接关系、验证关系和版本关系进行传播分析，识别直接影响对象和间接影响对象，并结合工程规则库判断是否存在接口冲突、需求冲突、验证项失效或仿真指标不满足。")
    add_para(doc, "预演结果以影响拓扑、影响矩阵和冲突清单形式展示。设计师确认后提交合并请求，审核人基于差异面板解决冲突。合并完成后，候选模型元素转正为发布分支中的权威模型元素，同时生成审计日志和版本记录。")

    add_heading(doc, "实施例六：仿真证据联动", 3)
    add_para(doc, "在低轨卫星互联网设计场景中，系统可以将轨道覆盖、链路可用度、载荷吞吐量、姿态控制动态响应等特定需求映射为仿真任务。对于轨道覆盖类需求，系统调用GMAT或等效轨道仿真工具生成覆盖率、可见窗口和轨道参数结果；对于动态响应类需求，系统调用Simulink或等效工具生成响应曲线、稳定时间和超调量结果。仿真结果被回写为验证证据，并与对应需求、参数、模型元素和校验项建立关联。")

    add_heading(doc, "可替代实施方式", 2)
    for text in [
        "知识增强检索可以采用向量数据库、关键词索引、图数据库、关系数据库或多种检索方式组合实现。",
        "工程知识图谱可以采用图数据库、关系型数据库、文档数据库或内存图结构实现。",
        "模型生成模块可以采用大语言模型、规则模板、模型转换器或多模型协同生成方式实现。",
        "模型校验规则可以来自模型语言规范、企业建模规范、接口控制规则、仿真约束规则或项目自定义规则。",
        "人在回路审核可以由设计师、知识工程师、系统管理员、领域专家或其他具备权限的角色完成。",
        "外部工具联动不限于MagicDraw、GMAT或Simulink，也可以替换为其他建模工具、仿真工具或数字工程平台。",
    ]:
        add_bullet(doc, text)


def build():
    doc = Document()
    set_styles(doc)
    set_header_footer(doc)

    add_para(doc, "发明专利申请文件初稿", size=23, bold=True, after=4)
    add_para(doc, "一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统", size=14, color=GRAY, after=14)

    add_table(
        doc,
        ["文件项目", "内容"],
        [
            ("申请类型", "发明专利"),
            ("拟申请名称", "一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统"),
            ("核心保护对象", "面向复杂系统工程模型的知识增强生成、自动校验、人在回路、变更影响分析和分支合并入库闭环。"),
            ("提交提示", "本文为申请文件初稿，正式提交前需由专利代理人结合检索结果进行权利要求稳定性、创造性和术语一致性复核。"),
        ],
        [1800, 7560],
        FILL,
    )

    add_heading(doc, "摘要", 1)
    add_para(doc, "本发明公开了一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统。该方法获取多源工程数据并抽取需求、功能、模块、接口、参数、约束和验证项等工程对象，构建包括向量语义索引、工程知识图谱和工程规则库的工程知识底座；响应于建模任务执行知识增强检索，生成具有来源证据链的候选工程模型；对候选工程模型执行语法、语义、接口、约束和仿真可满足性校验，并通过人在回路机制完成采纳、修正或驳回；针对模型变更在沙箱环境中进行影响分析和冲突识别，经审核后合并至发布分支并形成审计记录。本发明提高了工程模型生成效率和入库质量，实现AI生成内容的可追溯、可校验、可审核和可治理。")

    add_heading(doc, "核心创新点", 1)
    add_table(
        doc,
        ["创新点", "区别于常规AI问答或建模工具的技术特征"],
        [
            ("知识增强生成", "不是仅由大模型直接输出，而是由向量索引、工程知识图谱和规则库共同形成生成上下文。"),
            ("候选模型状态治理", "AI生成结果先进入候选状态，经过校验和人工审核后才能成为发布分支权威模型。"),
            ("校验与修复闭环", "对候选模型执行语法、语义、接口、约束和仿真可满足性校验，并形成可确认的修复补丁。"),
            ("沙箱变更预演", "变更先在沙箱中传播分析，输出影响拓扑、影响矩阵和冲突清单，再决定是否合并。"),
            ("证据链与审计", "需求、模型元素、仿真证据、审核记录和版本变更形成可追溯链路。"),
        ],
        [2100, 7260],
        BLUE_FILL,
    )

    add_claims(doc)
    add_description(doc)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
