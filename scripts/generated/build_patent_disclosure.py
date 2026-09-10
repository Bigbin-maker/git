from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\ai_mbse1\AI-MBSE-Demo")
OUT = ROOT / "docs" / "generated" / "一种面向复杂系统工程模型的知识增强式生成校验与变更治理方法及系统_专利交底书.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9DEE7"


def set_run_font(run, size=None, bold=None, color=None, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_para(doc, text="", style=None, size=11, bold=False, color=None, align=None, after=6, before=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12 if level == 2 else 8)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6 if level == 2 else 4)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=BLUE if level < 3 else DARK_BLUE)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="D7DEE8", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5)
    add_para(doc, "", after=4)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        set_cell_shading(hdr[i], header_fill)
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i == 0 and len(text) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_run_font(r, size=10)
    add_para(doc, "", after=4)
    return table


def set_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def set_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("专利交底书草案")
    set_run_font(r, size=9, color=MUTED)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("内部讨论稿 - 需由专利代理人结合检索结果进一步定稿")
    set_run_font(r, size=9, color=MUTED)


def build():
    doc = Document()
    set_styles(doc)
    set_header_footer(doc)

    add_para(doc, "专利交底书草案", size=23, bold=True, color=RGBColor(0, 0, 0), after=4)
    add_para(
        doc,
        "一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统",
        size=14,
        color=RGBColor(55, 55, 55),
        after=14,
    )
    meta = [
        ("文件性质", "发明专利申请前技术交底材料"),
        ("适用场景", "复杂系统工程设计、MBSE、SysML建模、工程知识库治理、模型变更管理"),
        ("核心主线", "工程数据接入 -> 知识增强生成 -> 自动校验修复 -> 人在回路审核 -> 变更影响分析 -> 冲突合并入库"),
        ("注意事项", "本文为技术与权利要求草案，不构成最终法律意见；正式提交前建议由专利代理人完成新颖性、创造性和权利要求稳定性检索。"),
    ]
    for k, v in meta:
        p = add_para(doc, after=2)
        r1 = p.add_run(f"{k}: ")
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(v)
        set_run_font(r2, size=11)

    add_callout(
        doc,
        "建议保护口径",
        "本申请不宜写成单纯的“AI生成SysML模型”或“RAG问答系统”，而应保护一种面向复杂系统工程模型的闭环治理方法：通过工程知识底座增强生成候选模型，再通过规则校验、人在回路、变更预演和分支合并机制，使AI生成结果能够被追溯、校验、审核并沉淀为权威工程模型。",
    )

    add_heading(doc, "一、拟申请名称", 1)
    add_para(doc, "一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统")

    add_heading(doc, "二、摘要", 1)
    add_para(
        doc,
        "本发明公开了一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统，涉及复杂系统设计、模型驱动工程和工程知识管理技术领域。该方法包括：获取需求文档、历史模型、接口规范、仿真结果、评审记录等多源工程数据；对多源工程数据进行工程对象抽取，形成包含需求、功能、模块、接口、参数、约束、验证项及其关系的工程知识底座；基于向量语义索引、工程知识图谱和规则库进行知识增强检索，生成候选工程模型及对应的可追溯证据链；对候选工程模型执行语法、语义、接口完整性、命名规范、约束一致性和仿真可满足性校验，并输出修复建议或自动规范化结果；通过人在回路机制对候选模型进行审核、采纳、驳回或修正；在模型发生变更时，于沙箱环境中进行变更影响预演，生成影响拓扑、影响矩阵和冲突提示；经审核通过后，将候选模型或变更结果合并至发布分支并形成审计记录。本发明能够提高复杂系统工程模型生成效率、降低模型入库错误率，并实现AI生成内容的可追溯、可校验、可审核和可治理。",
    )

    add_heading(doc, "三、创新点与知识产权保护重点", 1)
    add_table(
        doc,
        ["序号", "创新点", "要解决的问题", "技术手段", "技术效果"],
        [
            (
                "1",
                "面向工程模型的知识增强生成闭环",
                "通用大模型直接生成模型时容易脱离项目上下文，且缺乏工程可追溯性。",
                "将需求文档、模型元素、接口关系、规则库和历史评审记录组织为工程知识底座，并在生成前进行向量检索与图谱邻域检索。",
                "生成结果能够绑定来源证据，降低幻觉和上下文错配风险。",
            ),
            (
                "2",
                "工程对象级抽取与模型元素映射",
                "传统文档问答只能回答文本问题，难以直接转化为可治理的工程模型。",
                "从文档中抽取需求、功能、模块、接口、参数、约束、验证项等对象，并建立到工程模型元素的映射关系。",
                "使自然语言需求能够转化为可建模、可校验、可入库的结构化模型对象。",
            ),
            (
                "3",
                "生成-校验-修复一体化",
                "AI生成的模型即使语义合理，也可能违反SysML语法、项目命名规则或接口规范。",
                "在候选模型入库前执行语法规则、语义一致性、接口完整性、属性类型、连接合法性和项目规范校验，并生成修复动作。",
                "在入库前形成质量门禁，减少错误模型污染知识库。",
            ),
            (
                "4",
                "人在回路的候选模型治理",
                "完全自动入库缺乏工程责任边界，人工离线审核又无法沉淀反馈。",
                "将AI输出标记为候选模型，支持设计师采纳、修正、驳回和备注，并记录反馈与审计信息。",
                "兼顾AI效率和工程审核责任，并形成可用于后续优化的反馈数据。",
            ),
            (
                "5",
                "沙箱式变更影响预演",
                "工程模型参数或接口变更可能引起跨模块、跨需求、跨验证项连锁影响。",
                "在不覆盖发布分支的沙箱环境中模拟变更，基于图谱传播、规则约束和版本差异生成影响拓扑与影响矩阵。",
                "提交合并前即可识别直接影响、间接影响和潜在冲突。",
            ),
            (
                "6",
                "分支合并入库与权威模型转正",
                "候选模型、开发分支和发布分支之间缺少面向工程知识的治理机制。",
                "将候选模型变更提交为合并请求，审核人解决冲突后转正为发布分支中的权威模型元素。",
                "实现AI生成内容从候选状态到权威状态的可控流转。",
            ),
            (
                "7",
                "仿真证据与模型校验联动",
                "模型是否满足特定需求往往需要仿真工具验证，生成模型与仿真结果之间割裂。",
                "将GMAT、Simulink或其他仿真工具输出作为验证证据，与需求、模型元素和校验项建立关联。",
                "形成从需求到模型再到仿真证据的闭环验证链。",
            ),
        ],
        [720, 1720, 2240, 2520, 2160],
        BLUE_FILL,
    )

    add_heading(doc, "四、技术领域", 1)
    add_para(
        doc,
        "本发明涉及复杂系统工程、模型驱动工程、人工智能辅助设计、工程知识管理和模型变更治理技术领域，尤其涉及一种利用知识增强人工智能实现工程模型生成、模型校验、人在回路审核、变更影响分析和分支合并入库的方法及系统。",
    )

    add_heading(doc, "五、背景技术", 1)
    for text in [
        "复杂系统工程设计通常涉及需求、功能、逻辑架构、物理架构、接口、参数、约束、验证项、仿真结果和评审意见等多类工程数据。上述数据往往分散在需求文档、建模工具、图数据库、接口控制文档、仿真工具、评审记录和项目管理系统中，导致工程设计人员在模型构建、模型审查和模型变更时需要进行大量人工比对。",
        "现有大模型技术能够根据自然语言生成文本或代码，但在复杂系统工程模型场景中，直接使用大模型生成SysML、UML、AADL或其他工程模型存在三个问题：第一，生成过程缺乏项目知识约束，容易产生与已有需求或接口不一致的模型元素；第二，生成结果缺乏来源证据和状态标识，难以判断其是否可入库；第三，生成后的模型变更缺少影响分析和冲突治理，容易造成发布模型不一致。",
        "现有MBSE工具通常具备模型编辑、模型可视化和一定的规则检查能力，但其对自然语言需求解析、跨数据源知识检索、AI辅助建模、人工反馈沉淀、沙箱变更预演和分支合并治理的支持不足。尤其在多人协同和工程数据持续演进的场景下，缺少一种能够将知识增强生成、自动校验、人在回路审核、变更影响分析和入库治理贯通的通用方法。",
    ]:
        add_para(doc, text)

    add_heading(doc, "六、发明内容", 1)
    add_heading(doc, "6.1 发明目的", 2)
    add_para(
        doc,
        "本发明的目的在于提供一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统，用于解决AI生成工程模型缺乏工程上下文、缺乏可追溯校验、缺乏人工审核闭环以及缺乏变更治理机制的问题。",
    )

    add_heading(doc, "6.2 技术方案", 2)
    steps = [
        "S1，多源工程数据接入：获取需求文档、接口规范、历史工程模型、仿真结果、评审记录、版本记录和用户反馈数据。",
        "S2，工程对象抽取：从多源工程数据中识别需求、功能、场景、活动、模块、接口、参数、约束、验证项、仿真指标和角色权限等工程对象。",
        "S3，工程知识底座构建：建立工程对象之间的满足、派生、分配、连接、验证、依赖、冲突和版本关系，并构建向量语义索引、工程知识图谱和规则库。",
        "S4，知识增强检索：根据用户意图或建模任务，检索相关文档片段、模型元素、图谱邻域、历史变更和规则条目，形成生成上下文和证据链。",
        "S5，候选工程模型生成：基于生成上下文生成候选工程模型，候选工程模型包括文本化模型、图形化模型、模型元素清单、关系清单和来源证据。",
        "S6，模型校验与修复：对候选工程模型执行语法、语义、命名、接口、连接、约束、参数类型和仿真可满足性校验，并输出问题定位、风险等级和修复建议。",
        "S7，人在回路审核：将候选工程模型保持为未入库状态，由具备权限的设计师或审核人进行采纳、修正、驳回、备注或发起复核。",
        "S8，变更影响分析：针对模型元素、参数、接口或约束变更，在沙箱环境中进行预演，生成直接影响、间接影响、冲突项、受影响验证项和影响矩阵。",
        "S9，分支合并入库：将经审核的候选模型或变更提交为合并请求，由审核人解决冲突后合并至发布分支，形成权威工程模型。",
        "S10，反馈优化：将用户审核行为、校验结果、修复结果、合并结果和仿真证据回写至知识底座，用于后续检索、规则优化和生成策略优化。",
    ]
    for step in steps:
        add_number(doc, step)

    add_heading(doc, "6.3 系统组成", 2)
    add_table(
        doc,
        ["模块", "主要功能"],
        [
            ("数据接入模块", "接入需求文档、模型文件、接口规范、仿真结果、评审记录和历史版本数据。"),
            ("知识底座模块", "维护向量索引、工程知识图谱、规则库、模型元素库和证据链索引。"),
            ("任务理解模块", "识别用户意图、建模目标、当前项目上下文、附件状态和模型状态。"),
            ("模型生成模块", "基于知识增强上下文生成候选工程模型、模型元素关系和可视化图。"),
            ("模型校验模块", "执行语法校验、语义一致性校验、接口完整性校验、命名规范校验和仿真可满足性校验。"),
            ("人在回路模块", "提供采纳、修正、驳回、复核、备注和反馈记录能力。"),
            ("变更治理模块", "执行沙箱变更预演、影响传播分析、冲突检测和合并请求管理。"),
            ("入库与审计模块", "将通过审核的候选模型转正为发布分支权威元素，并记录操作者、时间、变更项和处理意见。"),
            ("工具联动模块", "与MagicDraw、GMAT、Simulink或其他工程工具进行模型同步、仿真调用和结果回传。"),
        ],
        [2200, 7160],
        LIGHT_FILL,
    )

    add_heading(doc, "6.4 有益效果", 2)
    for text in [
        "提高工程模型生成效率：设计人员能够从自然语言需求和已有工程知识直接生成候选模型，减少重复建模工作。",
        "增强工程可信度：生成结果绑定来源文档、图谱关系和规则条目，支持追溯和解释。",
        "降低模型入库风险：候选模型需经过自动校验和人工审核后才可入库，避免错误内容污染发布分支。",
        "提升变更治理能力：变更在沙箱中预演，提前识别直接影响、间接影响和潜在冲突。",
        "形成持续优化闭环：用户反馈、校验结果和合并结果可回写知识底座，用于后续生成策略和规则优化。",
        "具备跨模型语言扩展性：方法不限定于SysML，可扩展至UML、AADL、Simulink接口模型、数字孪生模型或企业自定义工程模型。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "七、附图说明", 1)
    for text in [
        "图1为本发明方法的总体流程图。",
        "图2为本发明系统的功能模块架构图。",
        "图3为工程知识底座的数据结构示意图。",
        "图4为知识增强候选模型生成流程图。",
        "图5为候选模型校验与人在回路审核流程图。",
        "图6为沙箱变更影响分析和分支合并入库流程图。",
        "图7为需求、模型元素、仿真证据和审计记录之间的追溯关系示意图。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "八、具体实施方式", 1)
    add_heading(doc, "8.1 总体实施例", 2)
    add_para(
        doc,
        "在一个实施例中，本发明应用于低轨通信卫星、无人系统、航空装备、工业生产线或其他复杂系统的工程设计过程。系统接收用户上传的需求说明、接口控制文件、已有SysML模型和仿真结果文件，将其解析为统一工程对象，并建立需求、功能、模块、接口、参数、约束、验证项之间的关系。用户发起建模任务时，系统先检索相关文档片段、图谱邻域和规则条目，再生成候选工程模型；候选模型经过自动校验和人工审核后，才可进入开发分支或发布分支。",
    )

    add_heading(doc, "8.2 知识增强生成机制", 2)
    add_para(
        doc,
        "系统在生成候选模型前，并不只依赖大模型的通用知识，而是根据当前项目、当前任务、用户角色、上传附件和已有模型状态构造检索请求。检索请求分别作用于向量语义索引和工程知识图谱：向量语义索引用于召回相似需求片段、接口说明和历史评审意见；工程知识图谱用于召回与目标需求相邻的模块、接口、约束、验证项和历史变更记录。检索结果经过去重、排序和证据标注后，作为模型生成上下文输入模型生成模块。",
    )
    add_para(
        doc,
        "模型生成模块输出的候选模型至少包括：模型元素标识、元素类型、元素名称、元素说明、来源证据、与其他元素的关系、候选状态和版本标识。对于SysML场景，候选模型可以表现为文本化SysML、需求图、用例图、内部模块图、活动图或接口关系图；对于其他场景，也可以输出UML、AADL、Simulink接口描述或企业自定义模型结构。",
    )

    add_heading(doc, "8.3 校验与修复机制", 2)
    add_para(
        doc,
        "模型校验模块对候选模型执行多层校验。第一层为语法校验，用于检查模型语言语法、元素定义和关系定义是否合法；第二层为工程语义校验，用于检查需求与模块、接口与连接、约束与参数、验证项与需求之间是否存在不一致；第三层为项目规范校验，用于检查命名、编号、属性类型、生命周期状态和权限边界是否符合项目规则；第四层为仿真可满足性校验，用于判断关键需求是否关联了仿真场景、仿真参数和仿真证据。",
    )
    add_para(
        doc,
        "当校验发现问题时，系统生成问题定位、风险等级、影响范围和修复建议。对于命名不规范、属性类型缺失、连接方向错误、接口引用缺失等可自动修复问题，系统可以生成修复补丁，并在用户确认后执行规范化修复。对于涉及工程判断的问题，系统仅给出建议，由用户在人工审核环节处理。",
    )

    add_heading(doc, "8.4 人在回路与状态治理", 2)
    add_para(
        doc,
        "候选模型生成后默认处于候选状态，不能直接成为权威发布模型。具有相应权限的设计师可以对候选模型执行采纳、局部修正、驳回或提交复核操作。系统记录用户身份、操作时间、被操作元素、修改前后差异、审核意见和反馈类型。该反馈不仅形成审计证据，还可作为后续提示词优化、规则库补充和模型微调的数据来源。",
    )

    add_heading(doc, "8.5 变更影响分析与合并入库机制", 2)
    add_para(
        doc,
        "当用户修改模型元素、接口、参数或约束时，系统不直接覆盖发布模型，而是在开发分支或沙箱中生成变更预演。变更治理模块基于工程知识图谱中的依赖关系、满足关系、连接关系、验证关系和版本关系进行传播分析，识别直接影响对象和间接影响对象，并结合规则库判断是否存在接口冲突、需求冲突、验证项失效或仿真指标不满足。",
    )
    add_para(
        doc,
        "预演结果以影响拓扑、影响矩阵和冲突清单形式展示。设计师确认后可提交合并请求，审核人基于可视化差异面板解决冲突。合并完成后，候选模型元素转正为发布分支中的权威模型元素，同时生成审计日志和版本记录。",
    )

    add_heading(doc, "8.6 仿真工具联动实施例", 2)
    add_para(
        doc,
        "在低轨卫星互联网设计场景中，系统可以将轨道覆盖、链路可用度、载荷吞吐量、姿态控制动态响应等特定需求映射为仿真任务。对于轨道覆盖类需求，系统调用GMAT或等效轨道仿真工具生成覆盖率、可见窗口和轨道参数结果；对于动态响应类需求，系统调用Simulink或等效工具生成响应曲线、稳定时间和超调量结果。仿真结果被回写为验证证据，并与对应需求、参数、模型元素和校验项建立关联。",
    )

    add_heading(doc, "九、权利要求书草案", 1)
    claims = [
        "1. 一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法，其特征在于，包括：获取复杂系统设计过程中的多源工程数据；对所述多源工程数据进行工程对象抽取，得到包括需求、功能、模块、接口、参数、约束和验证项在内的工程对象；基于所述工程对象构建工程知识底座，所述工程知识底座包括向量语义索引、工程知识图谱和工程规则库；响应于用户的建模任务，基于所述工程知识底座执行知识增强检索，得到与所述建模任务相关的上下文信息和来源证据；基于所述上下文信息生成候选工程模型；对所述候选工程模型执行模型校验，得到校验结果和修复建议；接收用户对所述候选工程模型的审核操作，使所述候选工程模型处于采纳、修正、驳回或待复核状态；响应于工程模型变更请求，在沙箱环境中执行变更影响分析，得到影响范围和冲突信息；在所述审核操作和所述变更影响分析满足预设条件后，将候选工程模型或变更结果合并至发布分支并生成审计记录。",
        "2. 根据权利要求1所述的方法，其特征在于，所述多源工程数据包括需求文档、接口控制文件、历史工程模型、建模工具导出文件、仿真结果、评审记录、版本记录和用户反馈数据中的一种或多种。",
        "3. 根据权利要求1所述的方法，其特征在于，所述工程对象抽取包括从自然语言文本中识别需求描述、功能动作、系统组成、外部参与者、接口名称、接口方向、参数名称、参数取值、约束条件、验证方式和生命周期状态，并为所述工程对象分配唯一标识和来源证据。",
        "4. 根据权利要求1所述的方法，其特征在于，所述工程知识图谱包括需求与模型元素之间的满足关系、模块之间的组成关系、接口之间的连接关系、参数与约束之间的约束关系、验证项与需求之间的验证关系、模型元素之间的依赖关系以及不同版本之间的变更关系。",
        "5. 根据权利要求1所述的方法，其特征在于，所述知识增强检索包括基于向量语义索引召回相关文档片段，基于工程知识图谱召回与目标工程对象相邻的模型元素和关系，并基于工程规则库召回与所述建模任务相关的建模规则或校验规则。",
        "6. 根据权利要求1所述的方法，其特征在于，所述候选工程模型包括文本化模型描述、图形化模型视图、模型元素清单、模型关系清单、接口连接清单、约束清单以及所述候选工程模型与来源证据之间的追溯关系。",
        "7. 根据权利要求1所述的方法，其特征在于，所述模型校验包括模型语言语法校验、模型语义一致性校验、接口完整性校验、命名规范校验、属性类型校验、连接合法性校验、约束一致性校验和仿真可满足性校验中的一种或多种。",
        "8. 根据权利要求1所述的方法，其特征在于，当所述校验结果表明候选工程模型存在可自动修复问题时，生成对应的模型修复补丁；在接收到用户确认操作后，将所述模型修复补丁应用于所述候选工程模型并重新执行模型校验。",
        "9. 根据权利要求1所述的方法，其特征在于，所述审核操作包括用户对候选工程模型的采纳、修正、驳回、备注、复核提交和权限审批，且所述审核操作被记录为反馈数据并回写至工程知识底座。",
        "10. 根据权利要求1所述的方法，其特征在于，所述变更影响分析包括在不改变发布分支的沙箱环境中模拟模型元素、接口、参数或约束的变更，基于工程知识图谱和工程规则库识别直接影响对象、间接影响对象、受影响验证项、潜在冲突项和建议处理措施。",
        "11. 根据权利要求1所述的方法，其特征在于，所述发布分支中的工程模型元素具有权威状态，所述候选工程模型在经审核和合并前具有候选状态，系统通过分支合并机制将满足预设条件的候选状态转换为权威状态。",
        "12. 根据权利要求1所述的方法，其特征在于，所述工程模型包括SysML模型、UML模型、AADL模型、Simulink接口模型、数字孪生模型或企业自定义工程模型中的一种或多种。",
        "13. 根据权利要求1所述的方法，其特征在于，还包括调用外部仿真工具执行与目标需求相关的仿真任务，接收仿真结果，并将所述仿真结果作为验证证据关联至对应需求、模型元素和校验项。",
        "14. 一种面向复杂系统工程模型的知识增强式生成、校验与变更治理系统，其特征在于，包括：数据接入模块，用于获取多源工程数据；工程对象抽取模块，用于从所述多源工程数据中抽取工程对象；知识底座模块，用于构建并维护向量语义索引、工程知识图谱和工程规则库；知识增强检索模块，用于根据建模任务检索上下文信息和来源证据；模型生成模块，用于生成候选工程模型；模型校验模块，用于对候选工程模型执行校验并生成修复建议；人在回路模块，用于接收用户对候选工程模型的审核操作；变更治理模块，用于在沙箱环境中执行变更影响分析并生成冲突信息；分支合并模块，用于将通过审核的候选工程模型或变更结果合并至发布分支；审计反馈模块，用于记录操作日志并将审核反馈回写至工程知识底座。",
        "15. 根据权利要求14所述的系统，其特征在于，所述系统还包括工具联动模块，所述工具联动模块用于与工程建模工具或仿真工具进行模型同步、仿真任务调用和结果回传。",
        "16. 一种电子设备，包括处理器和存储器，所述存储器中存储有计算机程序，所述计算机程序被所述处理器执行时实现权利要求1至13任一项所述的方法。",
        "17. 一种计算机可读存储介质，其上存储有计算机程序，所述计算机程序被处理器执行时实现权利要求1至13任一项所述的方法。",
    ]
    for claim in claims:
        add_para(doc, claim, after=8)

    add_heading(doc, "十、建议重点保护的权利要求组合", 1)
    add_table(
        doc,
        ["保护层级", "建议写法", "保护目的"],
        [
            ("主权利要求", "以“工程模型的知识增强生成、校验与变更治理方法”为核心，覆盖多源数据、知识底座、生成、校验、人在回路、变更预演和合并入库。", "保护完整闭环，避免被限定为某个工具或单一模型语言。"),
            ("从属权利要求", "分别限定RAG检索、工程知识图谱、规则库校验、候选状态、修复补丁、沙箱影响分析、仿真证据关联等特征。", "覆盖系统实际亮点，提高创造性支撑。"),
            ("系统权利要求", "按照模块写数据接入、知识底座、模型生成、校验、人在回路、变更治理、合并和审计。", "覆盖软件平台部署形态。"),
            ("设备/介质权利要求", "写电子设备和计算机可读存储介质。", "补充软件方法的常见保护形式。"),
        ],
        [1600, 5000, 2760],
        BLUE_FILL,
    )

    add_heading(doc, "十一、可替代实施方式", 1)
    for text in [
        "知识增强检索可以采用向量数据库、关键词索引、图数据库、关系数据库或多种检索方式组合实现。",
        "工程知识图谱可以采用图数据库、关系型数据库、文档数据库或内存图结构实现。",
        "模型生成模块可以采用大语言模型、规则模板、模型转换器或多模型协同生成方式实现。",
        "模型校验规则可以来自SysML语法、企业建模规范、接口控制规则、仿真约束规则或项目自定义规则。",
        "人在回路审核可以由设计师、知识工程师、系统管理员、领域专家或其他具备权限的角色完成。",
        "外部工具联动不限于MagicDraw、GMAT或Simulink，也可以替换为其他建模工具、仿真工具或数字工程平台。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "十二、后续提交前建议补充材料", 1)
    for text in [
        "补充至少三张附图：总体流程图、系统模块图、变更影响分析流程图。",
        "提供一个完整实施例，例如“低轨通信卫星宽带载荷分系统需求到SysML模型生成、校验和变更入库”的具体流程。",
        "整理系统截图作为交底材料辅助说明，但正式专利附图应转化为黑白线框流程图或模块图。",
        "请专利代理人进一步检索AI+MBSE、RAG生成模型、工程知识图谱、SysML自动生成、模型变更影响分析等方向的近似专利。",
        "正式提交前应避免在公开论文、网站、演示材料中披露过细的权利要求技术细节。",
    ]:
        add_bullet(doc, text)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
