# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\ai_mbse1")
OUT = ROOT / "专利申请" / "一种基于大模型智能体的工程模型协同生成、语义审查与工具编排方法及系统.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(82, 82, 82)
BLACK = RGBColor(0, 0, 0)
BORDER = "D9DEE7"
FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"


def set_run_font(run, size: float = 11, bold: bool = False, color=None, italic: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_format(p, before: float = 0, after: float = 6, line_spacing: float = 1.10) -> None:
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing


def init_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = ""
    r = header.add_run("发明专利申请文件初稿")
    set_run_font(r, size=9, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("内部草案 - 提交前请由专利代理人结合检索结果复核")
    set_run_font(r, size=9, color=GRAY)
    return doc


def add_para(
    doc: Document,
    text: str = "",
    size: float = 11,
    bold: bool = False,
    color=None,
    before: float = 0,
    after: float = 6,
    align=None,
    italic: bool = False,
    style: str | None = None,
):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, before=before, after=after)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    before = 16 if level == 1 else 12 if level == 2 else 8
    after = 8 if level == 1 else 6 if level == 2 else 4
    set_paragraph_format(p, before=before, after=after)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]], widths: list[int], fill=FILL) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        shade_cell(cell, fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(p, after=0, line_spacing=1.05)
        r = p.add_run(header)
        set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            set_paragraph_format(p, after=0, line_spacing=1.05)
            if i == 0 and len(text) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_run_font(r, size=10)
    add_para(doc, "", after=4)


def add_callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="CBD5E1")
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    set_paragraph_format(p, after=0, line_spacing=1.1)
    r = p.add_run(f"{label}：")
    set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    add_para(doc, "", after=3)


def add_claims(doc: Document) -> None:
    add_heading(doc, "权利要求书", 1)
    claims = [
        "1. 一种基于大模型智能体的工程模型协同生成、语义审查与工具编排方法，其特征在于，包括：获取用户自然语言输入、附件内容、当前工程模型上下文以及用户交互状态；对所述自然语言输入进行意图识别，确定与所述自然语言输入对应的工程建模任务类型；根据所述工程建模任务类型生成任务状态对象，所述任务状态对象至少包括项目标识、用户角色、当前选中模型元素、候选模型草案标识、可调用工具集合和任务阶段；基于所述任务状态对象，从工程文档库、模型关系图谱、候选模型草案、历史反馈记录和外部工具状态中检索并组织大模型提示上下文；调用大模型智能体基于所述大模型提示上下文生成结构化工程中间表示，所述结构化工程中间表示包括模型操作、模型元素、模型关系、模型视图、文本化模型描述、假设条件、证据引用和拟调用工具信息中的一种或多种；将所述结构化工程中间表示写入候选模型草案区，而不直接写入正式工程模型库；对所述候选模型草案执行格式校验、语义校验和关系一致性校验，得到校验结果；基于所述校验结果，由所述大模型智能体生成面向用户的语义解释和可执行修复建议；响应于用户确认操作，执行采纳、拒绝、重新生成、局部修复或工具调用操作；记录所述大模型智能体的提示上下文、响应内容、工具调用、草案差异和用户反馈，形成可审计交互链。",
        "2. 根据权利要求1所述的方法，其特征在于，所述工程建模任务类型包括直接问答、知识库问答、需求澄清、需求结构化、工程模型生成、内部块图细化、模型语义审查、变更影响分析、仿真验证、模型同步推送和审计查询中的一种或多种。",
        "3. 根据权利要求1所述的方法，其特征在于，所述意图识别包括：对用户自然语言输入进行关键词识别、上下文状态识别和大模型分类中的一种或多种；当识别置信度低于预设阈值时，将任务路由至需求澄清或直接问答；当用户输入包含生成、修改、推送、验证或变更分析意图时，将任务路由至相应的工程模型操作流程。",
        "4. 根据权利要求1所述的方法，其特征在于，所述大模型提示上下文包括带来源标识的文档片段、与目标模型元素相邻的模型图谱子图、候选模型草案的当前状态、用户历史修改反馈、可调用工具的能力描述和当前权限约束，并且按照任务类型、来源可信度、图谱距离、时间新近性和上下文长度预算进行排序、截断或摘要化处理。",
        "5. 根据权利要求1所述的方法，其特征在于，所述结构化工程中间表示采用预设模式约束，所述预设模式至少限定元素标识、元素名称、元素类型、元素描述、所属包路径、来源需求、关系源端、关系目标端、关系类型、视图类型、布局信息、证据引用、风险说明和修复动作字段。",
        "6. 根据权利要求5所述的方法，其特征在于，所述元素类型包括Requirement、Block、UseCase、Activity、Interface和ConstraintBlock中的一种或多种；所述关系类型包括contains、satisfy、trace、refine、allocate、verify和dependency中的一种或多种。",
        "7. 根据权利要求1所述的方法，其特征在于，所述候选模型草案区与正式工程模型库隔离，所述候选模型草案具有preview、revised、accepted或rejected状态；仅当候选模型草案通过校验并接收到具有权限的用户采纳操作后，才将对应模型元素、模型关系或模型视图提交至正式工程模型库或外部建模工具。",
        "8. 根据权利要求1所述的方法，其特征在于，所述格式校验、语义校验和关系一致性校验包括重复标识校验、元素类型校验、包路径校验、候选状态校验、关系端点存在性校验、关系类型校验、需求满足关系校验、接口分配关系校验、约束验证关系校验、属性类型和单位校验以及文本化模型与结构化模型同步校验中的一种或多种。",
        "9. 根据权利要求1所述的方法，其特征在于，所述可执行修复建议包括规范化包路径、规范化元素状态、删除悬空关系、规范化关系类型、新增satisfy关系、新增allocate关系、新增verify关系、补充属性类型、补充参数单位以及重新渲染文本化模型中的一种或多种。",
        "10. 根据权利要求1所述的方法，其特征在于，当用户选择一个Block类型模型元素作为焦点元素时，所述大模型智能体基于所述焦点元素的名称、描述、相邻需求、相邻接口、已有包含关系和用户原始意图，生成仅包含Block、Interface和connector关系的内部块图草案，并将该内部块图草案作为候选模型视图提交语义审查。",
        "11. 根据权利要求1所述的方法，其特征在于，所述工具调用操作包括调用工程知识检索工具、需求结构化工具、模型生成工具、模型校验工具、变更影响分析工具、仿真工具、SysML接口工具、MagicDraw或Cameo同步工具中的一种或多种；所述大模型智能体基于任务计划选择工具、生成工具调用参数、接收工具返回结果并更新任务状态对象。",
        "12. 根据权利要求11所述的方法，其特征在于，所述工具调用采用可审计调用链，所述可审计调用链记录工具名称、调用参数摘要、调用时间、调用结果摘要、异常信息、回退模式和与候选模型草案之间的关联关系。",
        "13. 根据权利要求1所述的方法，其特征在于，针对变更影响分析任务，先基于模型关系图谱和工程规则执行影响传播，得到直接影响对象、间接影响对象、影响路径、风险等级和处理建议，再由所述大模型智能体根据所述影响传播结果生成面向工程师的语义解释、评审清单和后续工具调用建议。",
        "14. 根据权利要求1所述的方法，其特征在于，所述用户反馈包括采纳、拒绝、元素编辑、关系编辑、一键修复、重新生成、推送确认和备注意见中的一种或多种；所述用户反馈被记录为反馈数据，并在后续大模型提示上下文构造、检索排序、修复建议生成或项目偏好约束中被引用。",
        "15. 根据权利要求1所述的方法，其特征在于，所述工程模型包括SysML模型、UML模型、AADL模型、Simulink接口模型、数字孪生模型、需求追溯模型或企业自定义工程模型中的一种或多种。",
        "16. 一种基于大模型智能体的工程模型协同生成、语义审查与工具编排系统，其特征在于，包括：交互输入模块，用于接收用户自然语言输入、附件内容和用户交互状态；意图识别与任务路由模块，用于确定工程建模任务类型并生成任务状态对象；上下文编排模块，用于从工程文档库、模型关系图谱、候选模型草案、历史反馈记录和外部工具状态中组织大模型提示上下文；大模型智能体模块，用于生成任务计划、结构化工程中间表示、语义解释和修复建议；草案管理模块，用于保存并管理候选模型草案；模型审查模块，用于执行格式校验、语义校验和关系一致性校验；修复执行模块，用于根据用户确认执行模型修复动作；工具编排模块，用于调用工程知识检索工具、建模工具、仿真工具或模型同步工具；审计反馈模块，用于记录提示上下文、响应内容、工具调用、草案差异和用户反馈。",
        "17. 根据权利要求16所述的系统，其特征在于，所述上下文编排模块还用于对检索到的文档片段、模型图谱子图、草案状态和历史反馈进行来源标识、可信度排序、长度压缩和证据引用编号。",
        "18. 一种电子设备，包括处理器和存储器，所述存储器中存储有计算机程序，其特征在于，所述计算机程序被所述处理器执行时实现权利要求1至15任一项所述的方法。",
        "19. 一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至15任一项所述的方法。",
    ]
    for claim in claims:
        add_para(doc, claim, after=8)


def add_description(doc: Document) -> None:
    add_heading(doc, "说明书", 1)
    add_heading(doc, "一种基于大模型智能体的工程模型协同生成、语义审查与工具编排方法及系统", 2)

    add_heading(doc, "技术领域", 2)
    add_para(
        doc,
        "本发明涉及复杂系统工程、模型驱动工程、人工智能辅助设计、大语言模型智能体、工程知识检索、工程模型校验和数字化工程工具集成技术领域，尤其涉及一种利用大模型智能体将工程师自然语言意图转化为可审查、可修复、可追溯且可推送的工程模型操作流的方法及系统。",
    )

    add_heading(doc, "背景技术", 2)
    for text in [
        "复杂系统工程设计通常需要在需求、功能、结构、接口、参数、约束、验证、仿真和交付模型之间建立一致关系。以MBSE为代表的模型驱动工程方法能够提高系统设计的一致性和可追溯性，但实际应用中仍依赖工程师手工完成需求澄清、模型分解、接口定义、图视图构建、校验问题解释和外部工具同步，建模门槛高、周期长且协同成本较大。",
        "大语言模型能够理解自然语言并生成文本、代码或结构化内容，因此可用于辅助工程建模。然而，若直接让大模型生成SysML或其他工程模型，容易出现输出格式不稳定、模型关系不完整、接口和约束缺失、幻觉内容难以识别、生成结果直接污染正式模型库、工具调用过程不可审计等问题。尤其在复杂系统工程场景中，模型不仅是说明文档，而是后续分析、仿真、评审和交付的依据，必须具备严格的结构约束、状态管理和可追溯记录。",
        "现有的RAG检索、工程知识图谱和规则校验技术能够为模型生成提供知识依据，但仍需要一种面向大模型智能体的任务编排机制，使大模型能够识别用户意图、选择合适工具、构造受控上下文、输出工程中间表示、解释校验结果、提出可执行修复建议，并在人工确认后完成修复、采纳、推送或仿真验证。现有技术中尚缺少将上述环节组织为统一、可审计、可控的大模型辅助工程模型操作流的方案。",
    ]:
        add_para(doc, text)

    add_heading(doc, "发明内容", 2)
    add_heading(doc, "要解决的技术问题", 3)
    for text in [
        "解决大模型直接生成工程模型时输出不可控、结构不稳定、模型关系和接口约束容易缺失的问题。",
        "解决大模型缺少工程任务路由能力，无法区分问答、澄清、生成、校验、变更、仿真和推送等不同工程任务的问题。",
        "解决大模型生成结果缺少草案隔离和人工确认机制，可能直接影响正式工程模型库的问题。",
        "解决模型校验结果难以被工程师快速理解，且校验问题与可执行修复动作之间缺少自动衔接的问题。",
        "解决大模型调用知识库、建模工具、仿真工具和同步工具时过程不可见、不可审计、难以复盘的问题。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "技术方案", 3)
    add_para(
        doc,
        "为解决上述问题，本发明提供一种基于大模型智能体的工程模型协同生成、语义审查与工具编排方法。该方法接收用户自然语言输入、附件内容和当前工程模型上下文，通过意图识别与任务路由确定任务类型；根据任务类型生成任务状态对象，并从工程文档库、模型关系图谱、候选模型草案、历史反馈记录和外部工具状态中构造大模型提示上下文；大模型智能体在预设模式约束下输出结构化工程中间表示；系统将该中间表示写入候选模型草案区，随后执行格式校验、语义校验和关系一致性校验；大模型智能体根据校验结果生成语义解释和可执行修复建议；用户确认后，系统执行修复、采纳、拒绝、重新生成或工具调用操作，并记录完整审计链。",
    )
    add_para(
        doc,
        "本发明还提供一种系统，包括交互输入模块、意图识别与任务路由模块、上下文编排模块、大模型智能体模块、草案管理模块、模型审查模块、修复执行模块、工具编排模块和审计反馈模块。各模块协同工作，使大模型不再是一次性文本生成器，而是受工程上下文、结构模式、工具能力、权限约束和人工反馈共同约束的工程模型智能体。",
    )

    add_heading(doc, "有益效果", 3)
    for text in [
        "通过任务路由机制，使大模型能够根据用户意图进入不同工程流程，避免普通问答误触发模型生成或推送。",
        "通过上下文编排机制，将文档、模型图谱、草案、反馈和工具状态转化为带证据的大模型提示上下文，降低幻觉和脱离项目上下文的风险。",
        "通过结构化工程中间表示和预设模式约束，使大模型输出可以被程序校验、预览、修复和推送。",
        "通过候选模型草案区隔离机制，使大模型生成内容必须经校验和人工确认后才能进入正式工程模型库。",
        "通过大模型语义审查和可执行修复建议，将规则校验问题转化为工程师容易理解和确认的操作建议。",
        "通过工具调用审计链，记录大模型智能体在知识检索、建模、校验、仿真和同步过程中的关键行为，提高系统可复盘性和责任可追溯性。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "附图说明", 2)
    for text in [
        "图1为本发明方法的总体流程图。",
        "图2为本发明系统的功能模块架构图。",
        "图3为意图识别与任务状态对象生成流程图。",
        "图4为大模型提示上下文编排流程图。",
        "图5为结构化模型草案生成、审查与修复流程图。",
        "图6为大模型智能体工具调用编排与审计流程图。",
        "图7为人机协同反馈闭环示意图。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "具体实施方式", 2)
    add_heading(doc, "实施例一：总体流程", 3)
    add_para(
        doc,
        "在一个实施例中，用户在工程建模平台中输入自然语言，例如“围绕低轨宽带卫星互联网系统生成总体需求、接口和内部块图”，或者“分析载荷发射功率上调对电源、热控和链路预算的影响”。系统接收该输入后，同时读取当前项目、用户角色、已选模型元素、候选草案、附件摘要、历史对话和工具可用状态，形成原始交互上下文。",
    )
    add_para(
        doc,
        "意图识别与任务路由模块对原始交互上下文进行识别。如果用户只是询问概念，则进入知识问答流程；如果用户要求生成模型，则进入模型生成流程；如果用户要求检查模型，则进入语义审查流程；如果用户描述参数或接口变更，则进入变更影响分析流程；如果用户要求运行轨道传播或动态响应验证，则进入仿真验证流程。由此避免将所有输入简单送入同一个大模型问答接口。",
    )
    add_table(
        doc,
        ["步骤", "处理内容", "输出"],
        [
            ("S1", "接收用户输入、附件内容、当前模型上下文和用户交互状态。", "原始交互上下文。"),
            ("S2", "识别任务类型并生成任务状态对象。", "任务类型、项目状态、用户角色和可调用工具集合。"),
            ("S3", "从文档库、模型图谱、草案、反馈和工具状态中组织提示上下文。", "带来源标识的大模型提示上下文。"),
            ("S4", "调用大模型智能体生成结构化工程中间表示。", "模型操作、元素、关系、视图、假设和证据引用。"),
            ("S5", "将中间表示写入候选模型草案区。", "可预览、可编辑的模型草案。"),
            ("S6", "执行格式、语义和关系一致性校验。", "校验报告、问题列表和风险说明。"),
            ("S7", "大模型智能体解释校验结果并生成修复建议。", "自然语言解释和可执行修复动作。"),
            ("S8", "用户确认后执行修复、采纳、拒绝、重新生成或工具调用。", "更新后的草案、正式模型或工具返回证据。"),
            ("S9", "记录提示、响应、工具调用、差异和反馈。", "可审计交互链。"),
        ],
        [900, 5300, 3160],
        BLUE_FILL,
    )

    add_heading(doc, "实施例二：意图识别与任务路由", 3)
    add_para(
        doc,
        "意图识别与任务路由模块可以采用规则、关键词、上下文状态和大模型分类相结合的方式。系统首先从用户输入中识别“生成、修改、推送、验证、仿真、分析影响、解释、查询”等动作词，再结合当前界面状态、用户选中对象和历史对话判断真实任务。例如，用户在模型预览区选中一个Block并输入“展开内部组成”，系统将其识别为内部块图细化任务；用户在变更分析页输入“带宽由500MHz改为800MHz”，系统将其识别为变更影响分析任务。",
    )
    add_para(
        doc,
        "路由结果不仅包括任务类型，还包括任务状态对象。任务状态对象可表示为包含project_id、workspace、role、selected_element_id、draft_id、task_stage、available_tools、permission_scope、conversation_digest和attachment_digest等字段的数据结构。大模型智能体后续只能基于该任务状态对象和工具能力描述生成可执行计划，从而避免越权推送或脱离当前项目上下文执行操作。",
    )
    add_table(
        doc,
        ["任务类型", "触发示例", "处理策略"],
        [
            ("直接问答", "你是什么模型、这个页面怎么理解", "调用大模型直接回答，不生成模型草案。"),
            ("知识问答", "接口约束有哪些依据", "检索文档和模型图谱后回答，并标明主要依据。"),
            ("模型生成", "生成一个完整SysML工程", "构造提示上下文，输出结构化模型草案。"),
            ("内部块图细化", "展开该Block内部组成", "以焦点Block为约束生成IBD候选视图。"),
            ("模型审查", "检查这个草案有什么问题", "运行规则校验并由大模型解释结果。"),
            ("变更分析", "将链路容量提高到800Mbps", "调用影响分析工具并生成评审建议。"),
            ("仿真验证", "运行轨道传播验证", "调用GMAT、Simulink或其他仿真适配器。"),
            ("模型推送", "推送到MagicDraw", "检查权限和草案状态后调用同步工具。"),
        ],
        [1500, 3100, 4760],
        FILL,
    )

    add_heading(doc, "实施例三：大模型提示上下文编排", 3)
    add_para(
        doc,
        "上下文编排模块根据任务状态对象执行多源检索。对文档型知识，系统召回与用户问题相关的需求说明、接口控制文件、设计规范和评审记录片段；对模型型知识，系统从模型关系图谱中召回与目标元素相邻的需求、Block、UseCase、Interface、ConstraintBlock和关系路径；对草案型知识，系统读取候选草案当前elements、relationships、diagram_views和sysml_text；对反馈型知识，系统读取用户历史采纳、拒绝、编辑和修复记录；对工具型知识，系统读取外部工具可用状态和调用约束。",
    )
    add_para(
        doc,
        "检索结果经去重、排序、压缩和来源标识后形成提示上下文。提示上下文不是简单拼接文本，而是按任务类型组织为任务目标、当前状态、已知事实、候选证据、输出模式、禁止事项和可调用工具说明。对于模型生成任务，提示上下文强调输出结构化元素和关系；对于审查任务，提示上下文强调问题定位和修复建议；对于变更任务，提示上下文强调影响路径和风险解释。",
    )
    add_callout(
        doc,
        "上下文编排要点",
        "本发明保护的重点之一在于把文档、图谱、草案、反馈和工具能力转化为大模型可使用且可审计的提示上下文，而不是简单把用户问题直接发送给大模型。",
    )

    add_heading(doc, "实施例四：结构化工程中间表示", 3)
    add_para(
        doc,
        "大模型智能体输出受预设模式约束的结构化工程中间表示。该中间表示可采用JSON、XML、SysML textual、数据库记录或其他结构化数据形式。为了便于程序校验和工具推送，中间表示通常包括project_name、summary、assumptions、elements、relationships、diagram_views、sysml_text、trace_links、tool_calls和repair_actions等字段。",
    )
    add_table(
        doc,
        ["字段", "说明", "作用"],
        [
            ("elements", "模型元素清单，包含id、name、type、description、package、source_requirement、status等。", "支持画布展示、入库和端点校验。"),
            ("relationships", "模型关系清单，包含source、target、type和description等。", "支持追溯、影响分析和工具同步。"),
            ("diagram_views", "图视图数据，包含视图类型、焦点元素、布局和视图内元素关系。", "支持需求图、用例图、内部块图等视图生成。"),
            ("sysml_text", "文本化工程模型描述。", "支持SysML textual预览和外部工具交换。"),
            ("assumptions", "大模型生成过程中引入的假设。", "供工程师审核和补充。"),
            ("evidence_refs", "来源文档、模型图谱或反馈记录引用。", "支持可追溯审查。"),
            ("repair_actions", "可执行修复动作及其参数。", "支持一键修复和再次校验。"),
        ],
        [1600, 4800, 2960],
        FILL,
    )
    add_para(
        doc,
        "系统对中间表示进行解析时，如果发现缺少必要字段、关系端点不存在、元素类型不在允许集合内或视图类型与元素类型不匹配，则不允许直接进入正式模型库，而是生成校验问题并返回给大模型智能体进行解释或重新生成。",
    )

    add_heading(doc, "实施例五：草案隔离、语义审查与一键修复", 3)
    add_para(
        doc,
        "候选模型草案区用于隔离大模型输出。草案区可以保存original_json、current_json、draft_status、source_prompt、generation_source和feedback_logs。用户可以在预览界面对元素名称、类型、描述、关系源端、关系目标端和关系类型进行编辑。每次编辑都记录修改前后差异，以便审计和回滚。",
    )
    add_para(
        doc,
        "模型审查模块对候选草案进行规则校验。若发现接口未分配到Block、需求未被Block satisfy、ConstraintBlock未verify需求、属性缺少value_type或单位、SysML文本与结构化模型不同步等问题，系统将问题列表提供给大模型智能体。大模型智能体将规则问题转化为面向工程师的自然语言解释，例如说明该问题为何影响追溯、接口一致性或验证完整性，并生成对应的可执行修复动作。用户确认后，修复执行模块才应用该动作。",
    )
    add_table(
        doc,
        ["校验问题", "大模型语义解释", "可执行修复动作"],
        [
            ("需求未被满足", "该需求缺少Block到Requirement的satisfy关系，后续无法证明设计满足该需求。", "新增Block -> Requirement的satisfy关系。"),
            ("接口未分配", "该接口未关联承载模块，接口责任边界不明确。", "新增Interface -> Block的allocate关系。"),
            ("约束未验证需求", "该约束没有指向被验证需求，无法形成验证闭环。", "新增ConstraintBlock -> Requirement的verify关系。"),
            ("属性缺少类型", "参数没有value_type或单位，后续难以参与计算和校验。", "补充PowerValue、BandwidthValue、DurationValue等类型和单位。"),
            ("文本视图不同步", "结构化模型已变化，但SysML textual未重新渲染。", "重新渲染sysml_text。"),
        ],
        [2100, 4700, 2560],
        BLUE_FILL,
    )

    add_heading(doc, "实施例六：内部块图动态细化", 3)
    add_para(
        doc,
        "当用户选中一个Block并请求展开内部结构时，大模型智能体不使用固定的输入-处理-输出模板，而是基于焦点Block的工程语义动态生成内部块图。例如，焦点Block为通信载荷时，内部结构可围绕信号接入、调制解调、波束控制、功放、遥测和热控接口展开；焦点Block为地面网关时，内部结构可围绕站控、路由、数据缓存、网络接口和运维监控展开。系统要求生成结果只包含Block、Interface和connector关系，并通过关系端点校验和视图布局校验后进入候选视图。",
    )
    add_para(
        doc,
        "该实施方式使大模型能够在已有顶层模型基础上进行局部深化，同时保持焦点范围、元素类型和关系类型受控，避免将需求、用例或外部Actor错误放入内部块图中。",
    )

    add_heading(doc, "实施例七：工具调用编排与审计", 3)
    add_para(
        doc,
        "大模型智能体可根据任务计划调用不同工具。对于知识问答任务，调用知识检索工具；对于模型生成任务，调用模型生成工具和草案管理工具；对于审查任务，调用模型校验工具；对于变更任务，调用影响分析工具；对于仿真任务，调用GMAT、Simulink或等效仿真适配器；对于推送任务，调用SysML API、MagicDraw、Cameo或其他建模工具同步接口。",
    )
    add_para(
        doc,
        "工具调用不是隐式执行，而是记录在可审计调用链中。每次调用记录工具名称、参数摘要、调用时间、调用结果摘要、异常信息、回退模式和关联草案。当前置工具不可用时，系统可进入本地交换包、模拟结果或只读解释模式，并在调用链中记录回退原因，避免用户误以为已完成真实推送或真实仿真。",
    )
    add_table(
        doc,
        ["工具类别", "典型工具", "审计记录"],
        [
            ("知识检索", "文档检索、模型图谱检索、规则库检索", "召回条数、来源、得分和摘要。"),
            ("模型生成", "模型生成服务、IBD生成服务", "输入上下文摘要、输出草案标识和生成来源。"),
            ("模型审查", "规则校验、语义审查", "问题数量、风险等级、修复动作。"),
            ("变更分析", "图谱传播、工程规则分析", "影响路径、直接影响、间接影响和建议。"),
            ("仿真验证", "GMAT、Simulink或等效工具", "脚本、结果路径、指标和证据关联。"),
            ("模型同步", "SysML API、MagicDraw、Cameo Bridge", "推送对象、成功状态、交换包路径或错误摘要。"),
        ],
        [1600, 3700, 4060],
        FILL,
    )

    add_heading(doc, "实施例八：在AI-MBSE系统中的应用", 3)
    add_para(
        doc,
        "在AI-MBSE系统的一个具体应用中，系统面向低轨宽带卫星互联网总体设计。设计师输入总体任务描述并上传需求文档后，系统先通过大模型智能体进行需求澄清和任务路由；当设计师点击生成时，系统构造包含需求片段、已确认需求、当前项目模型、历史反馈和工具状态的提示上下文，大模型智能体输出包含Requirement、Block、UseCase、Interface、Activity和ConstraintBlock的候选SysML模型。",
    )
    add_para(
        doc,
        "候选模型在预览区展示，设计师可修改元素或关系。模型审查模块识别缺失的satisfy、allocate、verify关系及属性类型问题，大模型智能体解释问题并给出一键修复建议。设计师确认后，草案被更新并重新校验。通过审核后，系统调用MagicDraw或Cameo同步接口生成建模工具可接收的交换数据；对于轨道传播、动态响应或链路可用性相关需求，系统调用仿真适配器生成验证证据，并将证据回写到对应需求和模型元素上。",
    )
    add_para(
        doc,
        "当设计师提出参数变更，例如提高星地链路带宽或修改载荷发射功率时，系统调用影响分析工具识别对载荷处理链路、电源热控、接口容量、网络调度和验证场景的影响；大模型智能体将影响矩阵解释为工程评审意见，提示需要复核的需求、接口、约束、仿真和交付物。所有提示、响应、工具调用、草案差异和用户反馈均进入审计反馈模块。",
    )

    add_heading(doc, "可替代实施方式", 2)
    for text in [
        "大模型可以是云端通用大语言模型、私有化部署大语言模型、领域微调模型或多个模型组成的协同智能体。",
        "工程文档库可以采用全文索引、向量数据库、关系数据库、文档数据库或多种存储方式组合实现。",
        "模型关系图谱可以采用图数据库、关系型数据库、内存图结构或建模工具原生API实现。",
        "结构化工程中间表示不限于JSON，也可以采用XML、YAML、SysML textual、数据库记录或建模工具交换格式。",
        "外部建模工具不限于MagicDraw或Cameo，也可以替换为其他SysML、UML、AADL、数字孪生或企业自定义建模平台。",
        "仿真工具不限于GMAT或Simulink，也可以替换为STK、Modelica、有限元工具、链路预算工具、离散事件仿真工具或企业自研仿真平台。",
        "用户反馈可以用于后续提示上下文构造，也可以进一步用于检索排序优化、规则库更新、模板更新或模型微调。",
    ]:
        add_bullet(doc, text)


def build() -> None:
    doc = init_doc()

    add_para(doc, "发明专利申请文件初稿", size=23, bold=True, color=BLACK, after=4)
    add_para(doc, "一种基于大模型智能体的工程模型协同生成、语义审查与工具编排方法及系统", size=14, color=GRAY, after=12)
    add_table(
        doc,
        ["文件项目", "内容"],
        [
            ("申请类型", "发明专利"),
            ("拟申请名称", "一种基于大模型智能体的工程模型协同生成、语义审查与工具编排方法及系统"),
            ("核心保护对象", "大模型智能体对工程建模任务的意图识别、上下文编排、受控结构化生成、语义审查、一键修复、工具调用与审计反馈机制。"),
            ("与第一篇区分", "第一篇侧重知识增强工程模型生成、校验与变更治理闭环；本篇侧重大模型智能体如何理解意图、组织上下文、输出受控中间表示、解释校验结果并编排工具。"),
            ("提交提示", "本文为申请文件初稿，正式提交前建议由专利代理人结合检索结果调整权利要求宽窄、术语一致性和创造性论证。"),
        ],
        [1900, 7460],
        FILL,
    )

    add_heading(doc, "摘要", 1)
    add_para(
        doc,
        "本发明公开了一种基于大模型智能体的工程模型协同生成、语义审查与工具编排方法及系统。该方法接收用户自然语言输入、附件内容和当前工程模型上下文，识别工程建模任务类型并生成任务状态对象；根据任务状态对象从工程文档库、模型关系图谱、候选模型草案、历史反馈记录和外部工具状态中构造带来源标识的大模型提示上下文；调用大模型智能体生成受预设模式约束的结构化工程中间表示，并将其写入候选模型草案区；对候选模型草案执行格式校验、语义校验和关系一致性校验；由大模型智能体根据校验结果生成语义解释和可执行修复建议；在用户确认后执行修复、采纳、拒绝、重新生成或工具调用，并记录提示上下文、响应内容、工具调用、草案差异和用户反馈。本发明使大模型从一次性文本生成转化为受工程上下文、结构模式、权限和工具能力约束的工程模型智能体，提高工程模型生成、审查、修复、仿真和同步过程的可控性与可追溯性。",
    )

    add_heading(doc, "核心创新点", 1)
    add_table(
        doc,
        ["创新点", "技术特征"],
        [
            ("大模型任务路由", "根据用户意图、界面状态、选中模型元素和权限确定问答、生成、审查、变更、仿真或推送流程。"),
            ("提示上下文编排", "将文档片段、模型图谱子图、草案状态、历史反馈和工具状态组织为带来源标识的上下文。"),
            ("受控中间表示", "要求大模型输出可解析、可校验的模型元素、关系、视图、文本化模型、证据和修复动作。"),
            ("草案隔离", "大模型结果先进入候选模型草案区，经校验和人工确认后才可进入正式模型库。"),
            ("语义审查与修复", "将规则校验问题解释为工程师可理解的风险说明，并生成可执行修复动作。"),
            ("工具调用审计", "记录知识检索、生成、校验、仿真和同步工具调用链，支持复盘和责任追溯。"),
        ],
        [2300, 7060],
        BLUE_FILL,
    )

    add_claims(doc)
    add_description(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
