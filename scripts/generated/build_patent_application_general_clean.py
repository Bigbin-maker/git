from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\ai_mbse1\AI-MBSE-Demo")
OUT = ROOT / "docs" / "generated" / "一种面向复杂系统工程模型的知识增强式生成校验与变更治理方法及系统_通用保护版.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(85, 85, 85)
BLACK = RGBColor(0, 0, 0)
BORDER = "D9DEE7"
FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def init_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    header = sec.header.paragraphs[0]
    header.text = ""
    run = header.add_run("发明专利申请文件优化稿")
    set_run_font(run, size=9, color=GRAY)

    footer = sec.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("通用保护版 - 提交前请由专利代理人复核")
    set_run_font(run, size=9, color=GRAY)
    return doc


def p(doc, text="", size=11, bold=False, color=None, before=0, after=6, align=None, style=None):
    para = doc.add_paragraph(style=style)
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.1
    if align is not None:
        para.alignment = align
    if text:
        run = para.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return para


def h(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_before = Pt(16 if level == 1 else 12 if level == 2 else 8)
    para.paragraph_format.space_after = Pt(8 if level == 1 else 6 if level == 2 else 4)
    run = para.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=BLUE if level < 3 else DARK_BLUE)


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.left_indent = Inches(0.5)
    para.paragraph_format.first_line_indent = Inches(-0.25)
    para.paragraph_format.space_after = Pt(5)
    para.paragraph_format.line_spacing = 1.167
    run = para.add_run(text)
    set_run_font(run, size=11)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = tbl_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tbl_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), BORDER)


def geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(cell)


def table(doc, headers, rows, widths, fill=FILL):
    tbl = doc.add_table(rows=1, cols=len(headers))
    geometry(tbl, widths)
    borders(tbl)
    for idx, header in enumerate(headers):
        c = tbl.cell(0, idx)
        shade(c, fill)
        pp = c.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = pp.add_run(header)
        set_run_font(run, size=10.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = tbl.add_row().cells
        for idx, text in enumerate(row):
            pp = cells[idx].paragraphs[0]
            pp.paragraph_format.space_after = Pt(0)
            pp.paragraph_format.line_spacing = 1.05
            run = pp.add_run(text)
            set_run_font(run, size=10)
    p(doc, "", after=4)


def build():
    doc = init_doc()

    p(doc, "发明专利申请文件优化稿", size=23, bold=True, color=BLACK, after=4)
    p(doc, "一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统", size=14, color=GRAY, after=12)
    table(
        doc,
        ["版本定位", "说明"],
        [
            ("通用保护版", "前半部分保护一种通用工程模型治理方法，不把权利要求限定在某一个系统或某一个建模工具上。"),
            ("实施例", "后半部分再以AI赋能MBSE系统作为具体实施例，承接RAG、SysML、MagicDraw、仿真、人在回路和分支入库等已实现能力。"),
            ("使用建议", "可作为提交专利代理人的主稿。正式递交前，建议代理人结合检索结果调整权利要求宽窄。"),
        ],
        [1800, 7560],
        FILL,
    )

    h(doc, "摘要", 1)
    p(doc, "本发明公开了一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统。该方法获取复杂系统设计过程中的多源工程数据，抽取需求、功能、模块、接口、参数、约束、验证项等工程对象，并构建包含语义索引、工程知识图谱和工程规则库的工程知识底座；响应于用户提出的建模或变更任务，基于工程知识底座检索相关知识并生成候选工程模型；在候选工程模型入库前，对其进行语法、语义、接口、约束和可验证性校验，并通过人在回路机制完成审核确认；当模型发生变更时，在沙箱环境中进行影响传播分析和冲突识别，经审核后再将候选模型或变更结果合并为发布分支中的权威工程模型。本发明使人工智能生成的工程模型能够被追溯、校验、审核、变更分析和受控入库，提高复杂系统设计过程中的模型生成效率和数据治理质量。")

    h(doc, "核心保护思路", 1)
    p(doc, "本申请的保护重点不是某一个前端界面、某一个大模型接口或某一个SysML图，而是一种面向复杂系统工程模型的闭环治理方法。该方法把知识增强生成、模型质量校验、人工审核、变更预演和入库治理组合为连续流程，使AI生成内容从“候选结果”转化为“权威工程模型”时具有明确的证据链、质量门禁和责任记录。")
    table(
        doc,
        ["保护点", "通用表述", "在本系统中的实施例"],
        [
            ("知识增强生成", "基于语义索引、工程知识图谱和规则库生成候选工程模型。", "RAG检索文档片段、图谱关系和SysML模型元素后生成模型。"),
            ("候选模型治理", "AI输出先处于候选状态，审核通过后才能入库。", "设计师在预览区采纳、修正或驳回生成的用例图、内部模块图。"),
            ("模型校验修复", "入库前执行模型规则校验并生成修复建议。", "检查SysML语法、接口完整性、命名规范、属性类型和连接合法性。"),
            ("变更影响分析", "模型变更先在沙箱环境中传播分析影响范围。", "对链路带宽、接口容量、载荷参数等变更生成影响拓扑和影响矩阵。"),
            ("证据与审计", "生成、审核、仿真、合并均形成可追溯记录。", "记录来源文档、模型元素、仿真结果、用户反馈和合并日志。"),
        ],
        [1700, 3900, 3760],
        BLUE_FILL,
    )

    h(doc, "权利要求书", 1)
    claims = [
        "1. 一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法，其特征在于，包括：获取复杂系统设计过程中的多源工程数据；从所述多源工程数据中抽取工程对象，所述工程对象至少包括需求、功能、模块、接口、参数、约束和验证项；基于所述工程对象构建工程知识底座，所述工程知识底座至少包括语义索引、工程知识图谱和工程规则库；响应于用户提出的建模任务或变更任务，基于所述工程知识底座进行知识增强检索，得到与所述任务相关的工程上下文和来源证据；基于所述工程上下文生成候选工程模型；对所述候选工程模型执行模型校验，得到校验结果和处理建议；通过人在回路机制接收用户对所述候选工程模型的审核操作；针对工程模型变更，在沙箱环境中进行影响分析和冲突识别；在所述候选工程模型或变更结果满足预设条件后，将其合并为权威工程模型，并生成审计记录。",
        "2. 根据权利要求1所述的方法，其特征在于，所述多源工程数据包括需求文档、接口规范、历史工程模型、仿真结果、评审记录、版本记录、权限数据和用户反馈数据中的一种或多种。",
        "3. 根据权利要求1所述的方法，其特征在于，所述工程知识图谱包括需求与模型元素之间的满足关系、模块之间的组成关系、接口之间的连接关系、参数与约束之间的约束关系、验证项与需求之间的验证关系以及模型元素之间的依赖关系。",
        "4. 根据权利要求1所述的方法，其特征在于，所述知识增强检索包括基于语义索引召回相关工程文本，基于工程知识图谱召回相关模型元素和关系，基于工程规则库召回相关建模规则或校验规则，并将召回结果组织为带有来源证据的工程上下文。",
        "5. 根据权利要求1所述的方法，其特征在于，所述候选工程模型包括文本化模型描述、图形化模型视图、模型元素清单、模型关系清单、接口连接清单、约束清单以及候选工程模型与来源证据之间的追溯关系。",
        "6. 根据权利要求1所述的方法，其特征在于，所述模型校验包括模型语言语法校验、工程语义一致性校验、接口完整性校验、命名规范校验、属性类型校验、连接合法性校验、约束一致性校验和需求覆盖性校验中的一种或多种。",
        "7. 根据权利要求1所述的方法，其特征在于，当所述校验结果表明候选工程模型存在可修复问题时，生成与所述问题对应的修复建议或模型修复补丁，并在接收到用户确认后更新候选工程模型。",
        "8. 根据权利要求1所述的方法，其特征在于，所述审核操作包括采纳、修正、驳回、提交复核、备注或审批，且所述审核操作被记录为反馈数据并回写至工程知识底座。",
        "9. 根据权利要求1所述的方法，其特征在于，所述影响分析包括在不改变权威工程模型的条件下模拟模型元素、接口、参数或约束的变更，并基于工程知识图谱和工程规则库识别直接影响对象、间接影响对象、受影响验证项和潜在冲突项。",
        "10. 根据权利要求1所述的方法，其特征在于，所述候选工程模型在审核和合并前具有候选状态，所述权威工程模型具有发布状态，系统通过状态转换和分支合并机制将候选状态转化为发布状态。",
        "11. 根据权利要求1所述的方法，其特征在于，还包括调用外部建模工具或仿真工具，并将外部建模工具或仿真工具返回的结果作为证据关联至需求、模型元素、参数、约束或验证项。",
        "12. 根据权利要求1所述的方法，其特征在于，所述工程模型包括SysML模型、UML模型、AADL模型、Simulink接口模型、数字孪生模型或企业自定义工程模型中的一种或多种。",
        "13. 一种面向复杂系统工程模型的知识增强式生成、校验与变更治理系统，其特征在于，包括：数据接入模块，用于获取多源工程数据；工程对象抽取模块，用于抽取工程对象；知识底座模块，用于构建并维护语义索引、工程知识图谱和工程规则库；知识增强检索模块，用于生成带有来源证据的工程上下文；模型生成模块，用于生成候选工程模型；模型校验模块，用于执行模型校验并生成处理建议；人在回路模块，用于接收审核操作；变更治理模块，用于执行沙箱影响分析和冲突识别；合并入库模块，用于将满足预设条件的候选工程模型或变更结果合并为权威工程模型；审计反馈模块，用于记录审计信息并回写反馈数据。",
        "14. 根据权利要求13所述的系统，其特征在于，还包括工具联动模块，用于与工程建模工具、仿真工具、图数据库、向量数据库或权限认证系统中的一种或多种进行数据交互。",
        "15. 一种电子设备，包括处理器和存储器，所述存储器存储有计算机程序，其特征在于，所述计算机程序被所述处理器执行时实现权利要求1至12任一项所述的方法。",
        "16. 一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至12任一项所述的方法。",
    ]
    for claim in claims:
        p(doc, claim, after=8)

    h(doc, "说明书", 1)
    h(doc, "技术领域", 2)
    p(doc, "本发明涉及复杂系统工程、模型驱动工程、人工智能辅助设计、工程知识管理和模型变更治理技术领域，尤其涉及一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统。")

    h(doc, "背景技术", 2)
    p(doc, "复杂系统设计通常涉及大量工程数据，包括需求、功能、模块、接口、参数、约束、验证项、仿真结果、评审记录和版本记录等。这些数据往往分散在文档、建模工具、仿真工具、数据库和项目管理系统中，导致工程人员在建模、校验和变更时需要反复检索和人工比对。")
    p(doc, "现有人工智能技术能够根据自然语言生成文本或代码，但在复杂系统工程模型场景中，直接生成模型容易出现来源不可追溯、接口关系不完整、约束不一致、模型元素状态不明确、变更影响不可见等问题。现有建模工具虽然可以支持模型编辑和部分规则检查，但通常没有将知识增强生成、候选状态治理、人在回路审核、沙箱变更预演和分支入库管理贯通起来。")
    p(doc, "因此，需要一种通用技术方案，使人工智能生成的工程模型不再停留于一次性输出，而能够进入可追溯、可校验、可审核、可变更分析和可合并入库的工程治理闭环。")

    h(doc, "发明内容", 2)
    p(doc, "本发明的目的在于提供一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统，用于解决AI生成工程模型缺乏工程上下文约束、缺乏质量门禁、缺乏人工审核闭环以及缺乏变更影响分析的问题。")
    p(doc, "为实现上述目的，本发明通过多源工程数据接入、工程对象抽取、工程知识底座构建、知识增强检索、候选工程模型生成、模型校验、人在回路审核、沙箱变更影响分析以及合并入库等步骤，形成从工程数据到权威工程模型的闭环治理流程。")
    p(doc, "本发明的有益效果在于：一是将工程知识作为生成约束，使模型生成结果具有来源证据；二是将生成结果设置为候选状态，避免AI输出直接污染权威模型；三是在入库前进行规则校验和人工审核，提高模型质量；四是在变更前进行沙箱预演，提前识别跨需求、跨接口、跨验证项的影响；五是通过审计和反馈机制沉淀用户修正经验，支持后续持续优化。")

    h(doc, "附图说明", 2)
    for item in [
        "图1为本发明方法的总体流程图。",
        "图2为本发明系统的功能模块架构图。",
        "图3为工程知识底座的数据结构示意图。",
        "图4为候选工程模型生成、校验与审核流程图。",
        "图5为沙箱变更影响分析与合并入库流程图。",
    ]:
        bullet(doc, item)

    h(doc, "具体实施方式", 2)
    h(doc, "实施例一：通用工程模型治理流程", 3)
    p(doc, "在一个通用实施例中，系统首先获取多源工程数据，并从中抽取工程对象。抽取结果不直接作为最终模型，而是进入工程知识底座。工程知识底座同时维护语义索引、工程知识图谱和工程规则库：语义索引用于召回相似文本和历史资料，工程知识图谱用于描述工程对象之间的依赖和追溯关系，工程规则库用于约束生成、校验和变更过程。")
    p(doc, "当用户提出建模任务时，系统先从工程知识底座检索相关上下文，再生成候选工程模型。候选工程模型经校验后进入人工审核环节。审核通过后，候选工程模型才能转换为权威工程模型；审核未通过时，候选工程模型可以被修正、驳回或重新生成。")

    h(doc, "实施例二：面向MBSE系统的具体实现", 3)
    p(doc, "在一个具体实施例中，上述方法应用于AI赋能MBSE系统。该系统接入需求文档、SysML模型、图数据库、向量数据库、MagicDraw模型同步结果、GMAT或Simulink仿真结果以及用户审核记录。系统通过RAG检索和工程知识图谱召回相关需求、接口、模块和约束，再生成SysML文本模型、用例图、内部模块图或接口关系图。")
    p(doc, "生成的SysML模型不直接入库，而是在前端预览区以候选模型形式展示。设计师可以采纳、修正或驳回。系统根据SysML语法、项目命名规范、接口完整性、属性类型和连接合法性进行预评审与校验，并对可修复问题给出修复建议。")
    p(doc, "当设计师提出参数或接口变更时，例如修改链路带宽、接口容量或载荷参数，系统在沙箱环境中基于工程知识图谱进行影响传播分析，输出受影响需求、模块、接口、验证项和仿真任务。审核人确认后，系统通过分支合并将候选模型或变更结果合并为发布分支中的权威模型元素，并记录操作日志。")

    h(doc, "实施例三：仿真证据关联", 3)
    p(doc, "对于需要仿真验证的需求，系统可以调用外部仿真工具生成验证证据。例如，在航天系统设计场景中，轨道覆盖类需求可以关联GMAT仿真结果，动态响应类需求可以关联Simulink仿真结果。仿真结果与需求、参数、模型元素和校验项建立追溯关系，用于支持模型校验和后续评审。")

    h(doc, "可替代实施方式", 2)
    for item in [
        "语义索引可以由向量数据库、全文检索引擎或二者组合实现。",
        "工程知识图谱可以由图数据库、关系型数据库、文档数据库或内存图结构实现。",
        "模型生成可以由大语言模型、规则模板、模型转换器或多模型协同机制实现。",
        "模型校验规则可以来自标准建模语言规范、企业建模规范、接口控制规则、仿真约束规则或项目自定义规则。",
        "工程模型不限于SysML，也可以包括UML、AADL、Simulink接口模型、数字孪生模型或企业自定义工程模型。",
    ]:
        bullet(doc, item)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
