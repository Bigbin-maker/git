from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "generated"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(24, 32, 51)
MUTED = RGBColor(90, 101, 121)
HEADER_FILL = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT_FILL = "F7FAFF"
BORDER = "D8E2EF"
GOLD_FILL = "FFF8E8"


def set_run_font(run, name="Calibri", east_asia="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def setup_doc(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = ""
    r = header.add_run(title)
    set_run_font(r, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("AI-MBSE 演示材料")
    set_run_font(r, size=9, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=23, color=INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(subtitle)
    set_run_font(r, size=12.5, color=MUTED)
    return doc


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        elem = tc_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tc_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), color)


def margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        elem = tc_mar.find(qn(f"w:{key}"))
        if elem is None:
            elem = OxmlElement(f"w:{key}")
            tc_mar.append(elem)
        elem.set(qn("w:w"), str(value))
        elem.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            margins(cell)
            borders(cell)


def para(doc, text="", bold=False, color=None, size=11, italic=False, after=6, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=size, color=color or INK, bold=bold, italic=italic)
    return p


def label_para(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(label)
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def callout(doc, title, lines, fill=CALLOUT_FILL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.18
        r = p.add_run(line)
        set_run_font(r, size=10.2, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, size=10, color=DARK_BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            if idx == 0:
                shade(cells[idx], "FBFCFE")
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.16
            r = p.add_run(str(text))
            set_run_font(r, size=9.5, color=INK, bold=(idx == 0))
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def segment(doc, title, ops, narration):
    doc.add_heading(title, level=1)
    doc.add_heading("画面操作", level=2)
    for item in ops:
        numbered(doc, item)
    doc.add_heading("讲解词", level=2)
    for item in narration:
        para(doc, item)


def build_polished_script():
    doc = setup_doc("AI-MBSE 系统功能演示讲稿（润色版）", "面向专家评审与演示视频录制 | 建议时长 10 分钟")
    callout(
        doc,
        "开场白",
        [
            "各位专家好。下面向大家演示我们的 AI 赋能 MBSE 系统。目前系统已覆盖甲方提出的主要功能需求，并将大模型问答、RAG 知识检索、SysML 模型生成、MagicDraw 联动、人工审核、模型校验、变更影响分析以及分支合并入库贯通在同一条工作流中。",
            "接下来，我将按照“登录系统、提问检索、智能生成模型、人工审核与校验、发起变更与影响分析、解决冲突并合并入库”的主线进行演示。",
        ],
    )
    matrix(
        doc,
        ["演示段落", "核心功能", "建议时长"],
        [
            ("1. 系统权限管理与交互界面", "角色权限、工作空间隔离、统一交互界面", "0:00-1:00"),
            ("2. 模型问答与 RAG 增强", "知识库上传、切片检索、可追溯问答", "1:00-2:30"),
            ("3. 需求解析、模型生成与人在回路", "需求解析、SysML 生成、图形预览、人工采纳", "2:30-4:30"),
            ("4. 模型预评审与校验", "语法校验、规范检查、一键修复", "4:30-6:00"),
            ("5. 变更影响分析", "沙箱预演、影响拓扑、影响矩阵", "6:00-8:00"),
            ("6. 冲突合并与数据入库管理", "分支合并、冲突解决、审计追踪", "8:00-10:00"),
        ],
        [2700, 5260, 1400],
    )
    segment(
        doc,
        "一、系统权限管理与交互界面",
        [
            "打开系统统一登录界面，展示系统支持按角色分配不同工作空间和功能权限。",
            "以“系统管理员”身份进入系统，展示项目智能体列表、项目隔离关系和系统状态标签。",
            "指出右上角关键组件状态：DeepSeek-V4 Pro 大模型、SysML v2、MagicDraw Bridge、仿真工具以及 SQLite 数据库接口。",
            "退出后切换为“设计师”角色，进入设计师工作空间。",
            "展示左侧项目目录、中间大模型对话区、右侧模型/图谱/代码/验证结果看板。",
        ],
        [
            "首先看到的是系统统一登录界面。系统可以按照不同角色分配工作空间和功能权限，保证项目数据和操作权限相互隔离。",
            "我先用系统管理员身份进入。可以看到，我们为每个项目都搭建了独立的智能体，项目之间工作空间相互隔离；系统管理员可以统一查看和管理全部项目。",
            "同时，右上角展示了系统关键组件的接入状态，包括 DeepSeek-V4 Pro 大模型、SysML v2 服务、MagicDraw Bridge、仿真工具以及数据库接口。",
            "接下来切换到设计师界面。左侧是项目目录，设计师只能在预先分配的项目路径下工作；中间是大模型对话区；右侧是模型、图谱、代码和验证结果看板。对话区支持 Markdown 流式渲染，系统会持续跟踪当前项目、上下文、附件和模型状态。",
        ],
    )
    segment(
        doc,
        "二、模型问答与 RAG 增强",
        [
            "以设计师身份进入 AI 智能体，先输入“在吗”，展示基础模型问答能力。",
            "向知识库上传《低轨通信卫星宽带载荷分系统关键需求与接口说明.docx》。",
            "提问：当前低轨通信卫星的宽带载荷分系统有哪些关键需求和接口？",
            "展示回答内容中出现了需求编号、接口编号、文档来源和模型元素引用。",
        ],
        [
            "接下来演示模型问答功能。我先向智能体发送“在吗”，可以看到系统能够直接响应设计师的问题。",
            "现在我向知识库上传一份名为《低轨通信卫星宽带载荷分系统关键需求与接口说明》的文档。上传后，系统会自动进行知识切片和索引构建。",
            "然后我提问：当前低轨通信卫星的宽带载荷分系统有哪些关键需求和接口？此时系统不再只依赖大模型的通用知识，而是通过 RAG 检索，把文档片段、图数据库中的模型关系、向量数据库中的语义片段以及 SysML 模型元素综合起来生成答案。",
            "可以看到，回答中不仅列出了关键需求和接口，还能追溯到来源文档或模型元素。这说明系统具备面向工程数据的可解释问答能力。",
        ],
    )
    segment(
        doc,
        "三、需求解析、模型生成与人在回路",
        [
            "发起需求：请你帮我生成一个由 100 颗卫星组成的低轨卫星互联网总体方案。",
            "根据系统提示上传《100颗低轨卫星互联网总体方案需求说明.docx》。",
            "等待系统解析文档中的需求、接口、模块、约束和行为线索。",
            "展示右侧 SysML 文本、用例图、内部模块图、接口关系与多视图预览。",
            "在预览与确认状态中检查模型，必要时调整连接线或接口关系，然后点击“采纳”。",
        ],
        [
            "现在进入模型生成环节。我先发起需求：请你帮我生成一个由 100 颗卫星组成的低轨卫星互联网总体方案。系统会分步骤引导设计师补充输入，并提示上传需求文档。",
            "上传需求文档后，系统会解析其中的需求、接口、模块、约束和行为线索，再生成 SysML 文本、模型元素和关系。",
            "右侧可以看到图形化预览：用例图展示用户链路和业务场景，内部模块图展示星座、卫星平台、载荷、星间链路、地面网关和运控系统之间的模块关系。",
            "这里是人在回路的第一处关键控制点。AI 生成的模型不会直接入库，而是停留在预览与确认状态。设计师可以检查模块、接口和连接线，必要时进行人工修正，然后点击采纳。系统会记录设计师的赞同或修正反馈，用于后续提示词优化和模型微调。",
        ],
    )
    segment(
        doc,
        "四、模型预评审与校验",
        [
            "针对刚生成的模型点击“执行模型校验”。",
            "展示预评审报告中的通过项、警告项、错误项和规范化建议。",
            "演示一键修复，例如补齐接口类型、属性类型或规范命名。",
            "再次执行校验，展示问题数量下降或状态转为通过。",
        ],
        [
            "模型生成之后，系统进入预评审和校验。平台可以按照 SysML 语法、项目命名规则、接口完整性、属性类型、连接合法性以及甲方建模规范进行自动检查。",
            "系统不仅指出问题，还会给出一键修复或规范化建议。例如当接口类型、属性类型或约束命名不完整时，系统可以辅助补齐并重新校验。",
            "这样模型在入库前就经过质量门禁，减少后续知识库污染和人工返工。",
        ],
    )
    segment(
        doc,
        "五、变更影响分析",
        [
            "在模型或图谱看板中选择“星间链路带宽”参数。",
            "发起模拟变更：将星间链路带宽调整到 50Gbps。",
            "点击“运行影响分析”或“变更预演”。",
            "展示影响拓扑图、影响矩阵、直接影响和间接影响。",
        ],
        [
            "接下来演示变更影响分析。假设我们将星间链路带宽调整到 50Gbps，如果直接入库，可能会影响多个分系统。",
            "系统会先在沙箱环境中进行变更预演，不改变发布分支中的权威数据。可以看到，直接影响包括星间通信载荷、路由交换单元、链路预算和星上处理带宽；间接影响包括电源功耗、热控散热、地面网关接入能力以及相关测试验证项。",
            "设计师可以在提交前直观看到变更影响的深度和广度，从而决定是否继续提交合并请求。",
        ],
    )
    segment(
        doc,
        "六、冲突合并与数据入库管理",
        [
            "设计师确认变更无误后，从开发分支提交合并请求。",
            "管理员或审核人进入合并界面，查看本次新增、修改或删除的模型元素和关系。",
            "展示系统自动识别出的冲突点和影响摘要。",
            "在冲突面板中完成冲突解决，点击“合并入库”。",
            "展示审计日志，确认操作人、时间、变更项、冲突处理方式和入库结果。",
        ],
        [
            "最后进入数据入库管理。设计师确认变更无误后，不会直接覆盖发布数据，而是从开发分支提交合并请求。",
            "管理员或审核人进入合并界面，可以看到本次提交新增、修改或删除了哪些模型元素和关系，也可以看到系统自动识别出的冲突点。",
            "审核人基于面板完成冲突解决后点击合并入库。此时，候选模型元素正式转正为发布分支中的权威图谱元素。",
            "系统同时记录审计日志，保留操作者、时间、变更项和处理意见，满足后续追踪和责任闭环。",
        ],
    )
    callout(
        doc,
        "结束语",
        [
            "通过这条演示链路可以看到，系统已经覆盖从设计知识检索、需求理解、模型生成、人工确认、预评审校验、变更影响分析到冲突合并入库的完整闭环。",
            "平台的价值不只是调用大模型，而是把大模型能力嵌入 MBSE 的数据治理流程中：所有 AI 生成内容都可追溯、可校验、可审核、可合并，并最终沉淀为发布分支中的权威设计知识。",
            "以上就是我们的功能演示，谢谢各位专家。",
        ],
    )
    out = OUT_DIR / "AI-MBSE系统功能演示讲稿_润色版.docx"
    doc.save(out)
    return out


def build_rag_doc():
    doc = setup_doc("低轨通信卫星宽带载荷分系统关键需求与接口说明", "知识库上传演示文档 | 用于 RAG 检索与模型问答")
    callout(
        doc,
        "文档用途",
        [
            "本文档用于 AI-MBSE 系统知识库上传演示。系统上传后可自动切片、向量化并建立需求、接口、模块和约束之间的可追溯关系。",
            "推荐演示提问：当前低轨通信卫星的宽带载荷分系统有哪些关键需求和接口？",
        ],
    )
    doc.add_heading("一、系统背景", level=1)
    para(doc, "宽带载荷分系统面向低轨通信卫星互联网任务，承担用户接入、星间转发、馈电链路、星上处理和网络路由等功能。该分系统需要在有限功耗、质量和热控条件下，实现高容量、低时延、可重构和高可靠的星上通信能力。")
    doc.add_heading("二、关键需求", level=1)
    matrix(
        doc,
        ["需求编号", "需求名称", "需求说明"],
        [
            ("REQ-BP-001", "多波束用户接入", "载荷应支持 Ku/Ka 频段多波束用户接入，单星可形成不少于 16 个可重构点波束，支持波束指向、带宽和功率的动态调整。"),
            ("REQ-BP-002", "宽带吞吐能力", "单星用户侧有效转发吞吐量不低于 20Gbps，网关馈电链路峰值吞吐量不低于 40Gbps，支持按业务优先级进行带宽调度。"),
            ("REQ-BP-003", "星间链路能力", "载荷应与星间链路终端协同工作，支持不少于 10Gbps 的星间数据转发能力，具备升级至 50Gbps 的接口余量。"),
            ("REQ-BP-004", "星上处理与路由", "载荷应支持星上数字透明转发、波束间交换、业务分类、路由表更新和网络管理指令下发。"),
            ("REQ-BP-005", "低时延服务", "星上处理引入的单跳处理时延应不超过 20ms，关键控制指令优先级应高于普通用户业务数据。"),
            ("REQ-BP-006", "可靠性与降级运行", "关键载荷单元应支持冗余备份、故障隔离和降级运行，单点故障不得导致整星宽带服务完全中断。"),
            ("REQ-BP-007", "遥测与健康管理", "载荷应输出功放温度、波束状态、链路误码率、吞吐量、队列拥塞和转发状态等遥测参数。"),
            ("REQ-BP-008", "在轨可重构", "载荷应支持在轨更新波束配置、路由策略、转发参数和业务优先级策略，更新过程应具备回滚机制。"),
        ],
        [1500, 1900, 5960],
    )
    doc.add_heading("三、外部接口", level=1)
    matrix(
        doc,
        ["接口编号", "接口名称", "连接对象", "接口说明"],
        [
            ("IF-BP-001", "用户链路接口", "用户终端", "提供 Ku/Ka 频段下行和上行接入，承载用户业务数据、同步信令和接入控制信息。"),
            ("IF-BP-002", "馈电链路接口", "地面网关站", "提供网关到卫星的高速馈电链路，用于用户业务回传、网络控制和路由策略更新。"),
            ("IF-BP-003", "星间链路接口", "相邻卫星", "承载跨星转发业务，支持链路状态交换、路由转发和拥塞控制信息交互。"),
            ("IF-BP-004", "平台电源接口", "卫星平台电源分系统", "为功放、数字处理器、路由交换单元和频率综合器供电，并提供功耗状态反馈。"),
            ("IF-BP-005", "热控接口", "卫星热控分系统", "输出功放、处理器和射频前端热耗参数，接收热控模式约束和告警状态。"),
            ("IF-BP-006", "测控管理接口", "星务计算机/运控系统", "接收配置、遥控和软件更新指令，输出载荷健康状态、告警和运行日志。"),
        ],
        [1350, 1700, 1800, 4510],
    )
    doc.add_heading("四、核心模块", level=1)
    for item in [
        "宽带通信载荷控制器：负责载荷任务调度、模式切换、配置下发和故障处理。",
        "多波束天线与射频前端：负责用户链路波束形成、频率转换、低噪声放大和功率放大。",
        "数字信道化处理单元：负责数字透明转发、信道化、交换、调制解调和带宽分配。",
        "星间链路适配单元：负责跨星业务封装、链路状态交换和转发策略适配。",
        "网络路由与资源管理单元：负责业务优先级、路由表、拥塞控制和服务质量管理。",
    ]:
        bullet(doc, item)
    doc.add_heading("五、约束与校验要点", level=1)
    for item in [
        "发射功率、功放效率和热耗应与电源分系统、热控分系统建立参数约束关系。",
        "用户链路接口、馈电链路接口和星间链路接口应在 SysML 模型中建立明确的 InterfaceBlock 或接口元素。",
        "关键需求应追溯到至少一个功能模块、一个外部接口和一个验证项。",
        "星间链路带宽从 10Gbps 提升到 50Gbps 时，应触发对星间链路适配单元、路由交换单元、电源和热控需求的影响分析。",
    ]:
        bullet(doc, item)
    out = OUT_DIR / "低轨通信卫星宽带载荷分系统关键需求与接口说明.docx"
    doc.save(out)
    return out


def build_requirement_doc():
    doc = setup_doc("100颗低轨卫星互联网总体方案需求说明", "模型生成演示文档 | 用于 SysML 用例图与内部模块图生成")
    callout(
        doc,
        "文档用途",
        [
            "本文档用于 AI-MBSE 系统模型生成演示。建议提示词：请根据这份文档，自动生成 100 颗低轨卫星互联网总体方案的 SysML 用例图和内部模块图（IBD）。",
            "文档内容包含任务目标、系统组成、关键需求、接口、约束和典型用例，便于系统抽取模型元素和关系。",
        ],
    )
    doc.add_heading("一、任务目标", level=1)
    para(doc, "建设由 100 颗低轨通信卫星组成的卫星互联网系统，面向区域宽带接入、应急通信、海洋航空连接和边远地区互联网覆盖等任务，提供低时延、高可靠、可扩展的天基通信服务。")
    doc.add_heading("二、总体架构", level=1)
    matrix(
        doc,
        ["组成部分", "建议模型元素", "说明"],
        [
            ("星座系统", "Block: LEOConstellation", "由 100 颗低轨通信卫星组成，建议采用 10 个轨道面、每轨道面 10 星的 Walker 类构型作为演示方案。"),
            ("通信卫星", "Block: CommunicationSatellite", "每颗卫星包含平台、宽带通信载荷、星间链路终端、星务计算机、电源、热控、姿轨控和测控单元。"),
            ("地面网关", "Block: GatewayStation", "负责馈电链路接入、核心网连接、网络控制和业务汇聚。"),
            ("用户终端", "Actor/Block: UserTerminal", "包括固定站、车载终端、船载终端、机载终端和便携应急终端。"),
            ("运控系统", "Block: OperationControlSystem", "负责星座运行管理、资源调度、健康监测、软件更新和任务规划。"),
            ("业务管理系统", "Block: ServiceManagementSystem", "负责用户接入、计费策略、服务质量、业务优先级和网络安全策略。"),
        ],
        [1800, 2350, 5210],
    )
    doc.add_heading("三、关键需求", level=1)
    matrix(
        doc,
        ["需求编号", "需求名称", "需求说明"],
        [
            ("REQ-LEO-001", "星座规模", "系统应由 100 颗低轨通信卫星组成，具备按批次扩展至更大星座规模的能力。"),
            ("REQ-LEO-002", "覆盖能力", "系统应面向指定区域提供连续宽带覆盖，重点保障海洋、航空、边远地区和应急场景。"),
            ("REQ-LEO-003", "用户容量", "系统应支持不少于 10000 个并发宽带用户接入，具备按业务优先级进行资源分配的能力。"),
            ("REQ-LEO-004", "星间互联", "卫星之间应具备星间链路能力，支持跨星路由、链路状态交换和业务转发。"),
            ("REQ-LEO-005", "端到端时延", "普通互联网业务端到端时延目标不高于 80ms，应急指挥业务应按高优先级保障。"),
            ("REQ-LEO-006", "链路带宽", "星间链路初始带宽不低于 10Gbps，系统设计应预留升级至 50Gbps 的接口、电源和热控余量。"),
            ("REQ-LEO-007", "可靠性", "系统应支持卫星故障隔离、链路绕转、网关切换和业务降级运行。"),
            ("REQ-LEO-008", "运维管理", "系统应支持星座健康监测、任务规划、在轨配置更新和运行日志审计。"),
        ],
        [1500, 1850, 6010],
    )
    doc.add_heading("四、主要接口", level=1)
    matrix(
        doc,
        ["接口编号", "接口名称", "源/目标", "接口说明"],
        [
            ("IF-LEO-001", "用户接入接口", "UserTerminal -> CommunicationSatellite", "承载用户业务数据、接入认证、同步信令和链路控制信息。"),
            ("IF-LEO-002", "馈电链路接口", "CommunicationSatellite -> GatewayStation", "承载业务回传、网关接入、网络控制和链路状态信息。"),
            ("IF-LEO-003", "星间链路接口", "CommunicationSatellite <-> CommunicationSatellite", "承载跨星路由、链路状态交换、业务转发和拥塞控制信息。"),
            ("IF-LEO-004", "运控管理接口", "OperationControlSystem -> Constellation", "承载任务规划、健康监测、软件更新、配置下发和告警回传。"),
            ("IF-LEO-005", "服务管理接口", "ServiceManagementSystem -> GatewayStation", "承载用户策略、服务质量配置、计费策略和业务优先级信息。"),
        ],
        [1350, 1650, 2500, 3860],
    )
    doc.add_heading("五、典型用例", level=1)
    matrix(
        doc,
        ["用例编号", "用例名称", "参与者", "说明"],
        [
            ("UC-001", "宽带用户接入", "用户终端、通信卫星、地面网关", "用户终端通过卫星接入互联网，系统完成认证、资源分配和业务转发。"),
            ("UC-002", "跨星业务转发", "通信卫星、相邻卫星、地面网关", "当用户业务不在当前网关覆盖路径内时，系统通过星间链路进行跨星转发。"),
            ("UC-003", "应急通信保障", "应急用户、运控系统、业务管理系统", "系统按高优先级保障应急用户，动态调整波束和路由资源。"),
            ("UC-004", "星座健康监测", "运控系统、通信卫星", "运控系统持续接收卫星遥测、链路状态和载荷告警，触发故障处置流程。"),
            ("UC-005", "链路带宽升级评估", "设计师、影响分析服务", "设计师将星间链路带宽调整至 50Gbps，系统执行沙箱预演和影响分析。"),
        ],
        [1400, 1700, 2500, 3760],
    )
    doc.add_heading("六、参数与约束", level=1)
    for item in [
        "星座规模：100 颗卫星，演示构型为 10 轨道面 x 10 星。",
        "轨道高度：建议演示值 1100km；轨道倾角建议演示值 53 度，可根据任务区域调整。",
        "星间链路带宽：初始不低于 10Gbps，升级目标为 50Gbps。",
        "单星用户转发能力：不低于 20Gbps；端到端时延目标不高于 80ms。",
        "关键参数变更必须触发影响分析，尤其是星间链路带宽、发射功率、载荷功耗、热耗和网关容量。",
    ]:
        bullet(doc, item)
    doc.add_heading("七、建议生成的 SysML 视图", level=1)
    for item in [
        "用例图：展示用户终端、地面网关、运控系统、设计师和典型业务用例之间的关系。",
        "内部模块图（IBD）：展示星座系统、通信卫星、宽带载荷、星间链路终端、地面网关、运控系统和业务管理系统之间的连接关系。",
        "需求图：展示 REQ-LEO-001 至 REQ-LEO-008 及其与模块、接口、用例的追溯关系。",
        "活动图：展示用户接入、跨星转发和变更影响分析流程。",
    ]:
        bullet(doc, item)
    out = OUT_DIR / "100颗低轨卫星互联网总体方案需求说明.docx"
    doc.save(out)
    return out


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = [build_polished_script(), build_rag_doc(), build_requirement_doc()]
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
