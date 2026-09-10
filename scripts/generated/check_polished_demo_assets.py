from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
FILES = [
    ROOT / "docs" / "generated" / "AI-MBSE系统功能演示讲稿_润色版.docx",
    ROOT / "docs" / "generated" / "低轨通信卫星宽带载荷分系统关键需求与接口说明.docx",
    ROOT / "docs" / "generated" / "100颗低轨卫星互联网总体方案需求说明.docx",
]

NEEDLES = {
    "AI-MBSE系统功能演示讲稿_润色版.docx": ["DeepSeek-V4 Pro", "RAG", "星间链路带宽", "合并入库"],
    "低轨通信卫星宽带载荷分系统关键需求与接口说明.docx": ["REQ-BP-001", "IF-BP-003", "50Gbps", "宽带载荷分系统"],
    "100颗低轨卫星互联网总体方案需求说明.docx": ["REQ-LEO-001", "IF-LEO-003", "UC-005", "100 颗低轨通信卫星"],
}


def main() -> None:
    for path in FILES:
        doc = Document(path)
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        tables = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        text = paragraphs + "\n" + tables
        missing = [item for item in NEEDLES[path.name] if item not in text]
        print(path.name)
        print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} missing={missing} ok={not missing}")


if __name__ == "__main__":
    main()
