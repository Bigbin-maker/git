from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "generated" / "AI-MBSE系统功能演示讲稿.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(24, 32, 51)
MUTED = RGBColor(90, 101, 121)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F7FAFF"
WHITE = "FFFFFF"
BORDER = "D8E2EF"


def set_run_font(run, name="Calibri", east_asia="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        elem = borders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        elem = tc_mar.find(qn(f"w:{name}"))
        if elem is None:
            elem = OxmlElement(f"w:{name}")
            tc_mar.append(elem)
        elem.set(qn("w:w"), str(value))
        elem.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)


def add_para(doc, text="", style=None, bold=False, color=None, size=None, italic=False, align=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_label_para(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(label)
    set_run_font(r1, size=11, color=DARK_BLUE, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=INK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    return p


def add_callout(doc, title, lines, fill=CALLOUT):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.18
        r = p.add_run(line)
        set_run_font(r, size=10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_matrix(doc, headers, rows, widths, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    hdr = table.rows[0].cells
    for idx, text in enumerate(headers):
        set_cell_shading(hdr[idx], header_fill)
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=10, color=DARK_BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            if idx == 0:
                set_cell_shading(cells[idx], "FBFCFE")
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.16
            r = p.add_run(text)
            set_run_font(r, size=9.6, color=INK, bold=(idx == 0))
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_section_title(doc, text):
    p = doc.add_heading(text, level=1)
    return p


def add_subtitle(doc, text):
    p = doc.add_heading(text, level=2)
    return p


def add_segment(doc, title, duration, functions, objective, operations, narration, highlights, fallback=None):
    add_section_title(doc, f"{title} ({duration})")
    add_label_para(doc, "演示功能：", functions)
    add_label_para(doc, "段落目标：", objective)
    add_subtitle(doc, "画面操作")
    for item in operations:
        add_numbered(doc, item)
    add_subtitle(doc, "建议讲解词")
    for para in narration:
        add_para(doc, para, after=6)
    add_subtitle(doc, "核心看点")
    for item in highlights:
        add_bullet(doc, item)
    if fallback:
        add_callout(doc, "演示口径提示", fallback, fill="FFF8E8")


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def set_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("AI-MBSE 系统功能演示讲稿")
    set_run_font(r, size=9, color=MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("内部演示材料")
    set_run_font(r, size=9, color=MUTED)


def build_doc():
    doc = Document()
    setup_styles(doc)
    set_header_footer(doc)

    add_para(doc, "AI-MBSE 系统功能演示讲稿", bold=True, color=INK, size=24, after=4)
    add_para(doc, "面向甲方演示视频录制 | 10 分钟建议版", color=MUTED, size=13, after=14)
    add_matrix(
        doc,
        ["项目", "内容"],
        [
            ("演示主线", "登录系统 -> 提问检索 -> 智能生成模型 -> 人工审核与校验 -> 发起变更与影响分析 -> 解决冲突合并入库"),
            ("覆盖能力", "数据管理、交互界面、模型问答、RAG 检索、模型生成、模型预评审与校验、变更影响分析、角色权限、人在回路"),
            ("建议角色", "系统管理员、设计师、知识工程师/审核人。当前环境若未单独配置知识工程师账号，可由系统管理员权限模拟审核人动作。"),
            ("录制方式", "建议使用浏览器全屏或 150% 缩放下录制，重点突出左侧导航、AI 对话区、右侧模型/图谱/证据面板。"),
        ],
        [1700, 7660],
        header_fill=LIGHT_GRAY,
    )
    add_callout(
        doc,
        "一句话开场",
        [
            "本次演示展示的是一个面向复杂系统工程的 AI-MBSE 协同设计平台。系统把大模型问答、RAG 知识检索、SysML 模型生成、MagicDraw 联动、人工审核、模型校验、变更影响分析和分支合并入库贯通在同一条工作流里。",
        ],
    )

    add_section_title(doc, "一、演示节奏总览")
    add_matrix(
        doc,
        ["时间", "演示段落", "关键画面", "证明的能力"],
        [
            ("0:00-1:00", "登录与交互界面体验", "统一登录页、角色切换、AI 智能体工作台", "用户角色与权限管理、交互界面、工作空间隔离"),
            ("1:00-2:30", "知识检索与模型问答", "对话区提问、RAG 来源、右侧文档/模型定位", "模型问答、数据追溯、知识图谱与向量库智能路由"),
            ("2:30-4:30", "需求解析与模型生成", "上传 docx、生成 SysML、预览用例图/IBD、采纳", "模型生成、人在回路、候选模型管理"),
            ("4:30-6:00", "模型预评审与智能校验", "执行校验、预评审报告、一键修复", "模型预评审、规范校验、质量闭环"),
            ("6:00-8:00", "变更预演与影响分析", "参数变更、沙箱预演、影响拓扑和矩阵", "变更影响分析、直接/间接影响追踪"),
            ("8:00-10:00", "冲突合并与数据入库", "合并请求、冲突解决、审计日志、发布分支", "设计数据管理、分支合并、人在回路、权威知识库入库"),
        ],
        [1300, 2050, 2850, 3160],
    )

    add_section_title(doc, "二、录制前准备")
    for item in [
        "确认后端服务、前端服务、SQLite 数据库、LLM、SysML 服务、MagicDraw Bridge、仿真工具状态均已启动；首页顶部状态标签尽量保持 Real/Bridge/SQLite 正常。",
        "准备一份示例文档：《宽带通信载荷需求更新说明.docx》。如果没有正式文档，可使用系统已有技术文档或演示需求文本代替。",
        "提前准备提问语句：当前低轨通信卫星的宽带载荷分系统有哪些关键需求和接口？",
        "提前准备变更语句：将天线发射功率由 35dBW 变更为 40dBW。",
        "录屏时尽量保持左侧导航可见，便于甲方理解系统不是单点能力，而是覆盖完整工程流程。",
    ]:
        add_bullet(doc, item)

    add_segment(
        doc,
        "三、登录与交互界面体验",
        "0:00-1:00",
        "用户角色与权限管理、交互界面",
        "用最短时间证明系统具备企业级入口、角色隔离和面向设计工作的统一交互台。",
        [
            "打开系统统一登录界面，先选择“系统管理员”。",
            "简要展示管理员可见的全局状态、角色/工作空间入口和系统导航。",
            "退出或切换为“设计师”角色，进入设计师的开发工作空间。",
            "进入 AI 智能体或 AI 工作台页面，展示左侧项目区、中间大模型对话区、右侧 SysML/图谱/证据看板。",
            "在聊天区展示 Markdown 流式渲染、上下文状态、附件上传和快捷操作按钮。",
        ],
        [
            "首先进入系统的统一登录入口。这里可以对接甲方已有的 LDAP 或 OAuth2 认证体系，并在登录后按角色分配不同的工作空间和功能权限。",
            "我先用系统管理员身份进入，可以看到系统级状态、数据服务状态以及全局功能导航。随后切换为设计师角色，进入分配给设计师的开发分支工作空间。",
            "设计师的主要操作集中在 AI 工作台。左侧是项目与数据入口，中间是大模型对话区，右侧是模型、图谱、代码和验证结果的看板。对话区支持 Markdown 流式渲染，系统会持续跟踪当前项目、上下文、附件和模型状态。",
        ],
        [
            "角色权限不是静态页面展示，而是影响可见菜单、工作空间和后续入库审批动作。",
            "交互界面将“问答、生成、预览、校验、推送、合并”集中在同一个工作流中。",
            "顶部状态标签可辅助说明系统已接入 LLM、SysML、MagicDraw Bridge、数据库与仿真能力。",
        ],
        [
            "如果当前演示环境只配置了“系统管理员”和“设计师”两个账号，可说明知识工程师/审核人是同一权限框架下的可扩展角色，本次用管理员权限模拟审核人动作。",
        ],
    )

    add_segment(
        doc,
        "四、知识检索与模型问答",
        "1:00-2:30",
        "模型问答、设计数据管理、RAG 智能路由",
        "证明系统不是普通聊天机器人，而是能基于设计文档、知识图谱、向量数据库和 SysML 元素进行可追溯问答。",
        [
            "在设计师工作台对话框输入：当前低轨通信卫星的宽带载荷分系统有哪些关键需求和接口？",
            "等待模型流式输出答案，画面保留回答中的需求、接口、约束和建议。",
            "展示回答中的来源引用、文档片段或模型元素标识。",
            "点击引用或相关元素，让右侧看板高亮对应的需求段落、SysML 元素或接口节点。",
            "说明系统会根据问题类型自动路由到知识库、图数据库、向量数据库或模型仓库。",
        ],
        [
            "接下来进行模型问答。我输入一个面向设计人员的问题：当前低轨通信卫星的宽带载荷分系统有哪些关键需求和接口？",
            "系统不会只依赖大模型的通用知识，而是通过 RAG 检索把图数据库中的模型关系、向量数据库中的文档片段以及 SysML 模型元素综合起来，再生成答案。",
            "回答中的每个关键结论都可以追溯到来源。点击这里的引用或元素标识，右侧面板会定位到相应的文档段落或模型元素，例如需求、接口、约束或模块节点。这证明系统具备设计数据可追溯能力。",
        ],
        [
            "突出“可解释问答”：答案不是孤立文本，而是带来源、可定位、可追踪。",
            "突出“智能路由”：同一个问题可以访问文档库、模型库、知识图谱和向量库。",
            "为后续模型生成做铺垫：问答结果可以直接成为生成 SysML 模型的上下文。",
        ],
        [
            "如果录制时没有语音输入环境，可以直接使用文本输入；讲解时保留“系统支持语音或文本输入”的表述即可。",
        ],
    )

    add_segment(
        doc,
        "五、需求解析与模型生成",
        "2:30-4:30",
        "模型生成、人在回路第一环节",
        "展示从非结构化需求文档到 SysML 模型候选方案的自动化生成，并强调生成结果需人工确认后才能进入后续流程。",
        [
            "上传《宽带通信载荷需求更新说明.docx》。",
            "输入提示词：请根据这份文档，自动生成宽带通信载荷分系统的 SysML 用例图和内部模块图（IBD）。",
            "等待系统生成 SysML 文本、模型元素、关系和图表预览。",
            "展示多套备选生成方案或不同视图，例如用例图、内部模块图、活动/状态/序列相关视图。",
            "双击或点击图中节点，展示元素下钻、接口关联或 IBD 下钻能力。",
            "在预览与确认界面点击“采纳”，或演示手动调整一根连接线/接口关系后再采纳。",
        ],
        [
            "现在进入模型生成环节。我上传一份新的需求更新说明，并要求系统自动生成宽带通信载荷分系统的 SysML 用例图和内部模块图。",
            "系统会先解析文档中的需求、接口、模块、约束和行为线索，再生成 SysML 文本、模型元素和关系。右侧可以看到图形化预览，用例图展示用户链路和业务场景，内部模块图展示模块、接口和连接关系。",
            "这里是人在回路的第一处关键控制点。AI 生成的模型不会直接入库，而是停留在预览与确认状态。设计师可以检查图中的模块、接口和连接线，必要时做人工修正，然后点击采纳。系统会记录这次赞同或修正反馈，后续可用于提示词优化和模型微调。",
        ],
        [
            "强调“候选模型”概念：AI 负责提效，人负责确认。",
            "展示图形预览和 SysML 文本同步，证明不是只生成图片，而是生成可管理的模型资产。",
            "展示下钻查看接口/模块关系，体现模型完整性和可解释性。",
        ],
    )

    add_segment(
        doc,
        "六、模型预评审与智能化校验",
        "4:30-6:00",
        "模型预评审与校验",
        "证明平台可以在模型入库前按 SysML 语法和甲方建模规范做自动检查，并给出可执行修复建议。",
        [
            "针对刚生成的内部模块图或候选模型，点击“执行模型校验”或进入预评审/校验区域。",
            "等待系统输出预评审报告，展示通过项、警告项、错误项和规范建议。",
            "重点展示一条典型问题：警告：天线发射功率属性未定义类型（Untyped value types）。",
            "点击“一键修复/规范化”按钮，展示系统补齐类型、约束或命名规范。",
            "再次执行校验，展示问题数下降或状态变为通过/可发布。",
        ],
        [
            "模型生成之后，系统会进入预评审和校验。这里可以按 SysML 语法、项目命名规则、接口完整性、属性类型、连接合法性以及甲方建模规范进行自动检查。",
            "例如系统识别到天线发射功率属性没有定义明确类型，这会影响后续参数约束、仿真和变更分析。系统不仅指出问题，还给出一键修复或规范化建议。",
            "点击修复后，系统会自动补齐类型或规范命名，再重新校验。这样模型在入库前就经过质量门禁，减少后续知识库污染和人工返工。",
        ],
        [
            "预评审不是最终审批，而是面向设计人员的前置质量门禁。",
            "错误提示要尽量停留 3-5 秒，让观众看清“问题、原因、修复建议”。",
            "修复后再校验一次，形成“发现问题 -> 修复 -> 复核”的闭环。",
        ],
    )

    add_segment(
        doc,
        "七、变更预演与影响分析",
        "6:00-8:00",
        "变更影响分析",
        "证明系统能够在正式修改前用沙箱做变更预演，识别直接影响和间接影响。",
        [
            "在图谱或模型看板中选中“天线发射功率”参数。",
            "发起模拟变更：将发射功率由 35dBW 变更为 40dBW。",
            "点击“运行影响分析”或“变更预演”。",
            "展示沙箱分析结果：影响拓扑图、影响矩阵、风险等级和建议动作。",
            "重点讲解直接影响：电源子系统功耗、放大器余量、链路预算。",
            "继续讲解间接影响：热控散热需求、结构布置、供电约束、测试验证项。",
        ],
        [
            "接下来演示变更影响分析。假设设计师需要把天线发射功率从 35dBW 提升到 40dBW，如果直接改入库，可能会影响多个分系统。",
            "系统会先在沙箱环境中进行变更预演，不改变发布分支的权威数据。分析结果会生成影响拓扑图和影响矩阵。",
            "可以看到，直接影响包括电源子系统需要提供更高功率、链路预算需要重新计算；间接影响包括热控分系统散热需求增加，以及相关测试验证项需要同步更新。设计师可以在提交前直观看到这次变更的深度和广度。",
        ],
        [
            "突出“沙箱预演”：先分析、再提交，避免破坏正式知识库。",
            "拓扑图适合讲影响范围，矩阵适合讲影响类型、风险等级和处理建议。",
            "把变更影响和前面的模型关系、接口关系、需求追溯连接起来讲。",
        ],
    )

    add_segment(
        doc,
        "八、冲突合并与数据入库管理",
        "8:00-10:00",
        "设计数据管理、分支管理与冲突合并、人在回路最终环节",
        "展示候选设计从开发分支经过审核、冲突解决和审计记录后转正为发布分支权威知识。",
        [
            "设计师确认变更分析结果后，创建从开发分支到发布分支的合并请求（Merge Request）。",
            "切换为知识工程师/审核人角色；若当前环境未单独配置该角色，可用系统管理员权限模拟审核动作。",
            "打开合并请求，展示实体变更列表、关系变更列表、冲突点和影响摘要。",
            "在冲突面板中选择保留版本、采用候选版本或手动合并。",
            "点击“合并入库”，展示发布分支更新成功。",
            "展示审计日志：操作者、时间、变更内容、冲突处理方式和入库结果。",
        ],
        [
            "最后进入数据入库管理。设计师确认变更无误后，不会直接覆盖发布数据，而是从开发分支提交合并请求。",
            "审核人进入合并界面，可以看到这次提交新增、修改或删除了哪些模型元素和关系，也可以看到系统自动识别出的冲突点。",
            "审核人基于面板完成冲突解决后点击合并入库。此时，候选模型元素正式转正为发布分支中的权威图谱元素。系统同时记录审计日志，保留操作者、时间、变更项和处理意见，满足后续追踪和责任闭环。",
        ],
        [
            "强调“人在回路最终环节”：AI 和设计师都不能绕过审核直接污染发布库。",
            "冲突解决面板是数据治理核心画面，建议录制时停留较长时间。",
            "审计日志是面向甲方管理要求的重要证明。",
        ],
    )

    add_section_title(doc, "九、功能覆盖矩阵")
    add_matrix(
        doc,
        ["甲方要求能力", "演示位置", "建议讲解重点"],
        [
            ("数据管理", "第 2、6 段", "文档、模型、图谱、分支、合并、审计统一管理。"),
            ("交互界面", "第 1 段", "AI 对话区 + 数据管理 Web 看板 + 模型/图谱/证据联动。"),
            ("模型问答", "第 2 段", "基于 RAG、知识图谱、向量库和 SysML 元素的可追溯问答。"),
            ("模型生成", "第 3 段", "从需求文档生成 SysML 文本、元素、关系、用例图和 IBD。"),
            ("变更影响分析", "第 5 段", "沙箱预演、影响拓扑、影响矩阵、直接/间接影响。"),
            ("模型预评审与校验", "第 4 段", "语法检查、规范检查、预评审报告、一键修复。"),
            ("用户角色与权限管理", "第 1、6 段", "管理员、设计师、知识工程师/审核人及工作空间权限。"),
            ("人在回路", "第 3、4、6 段", "采纳/修正、校验确认、冲突审核、合并入库。"),
            ("RAG 等智能能力", "第 2、3 段", "检索增强生成、来源定位、上下文追踪、模型生成增强。"),
        ],
        [2100, 1900, 5360],
    )

    add_section_title(doc, "十、收尾总结话术")
    add_callout(
        doc,
        "建议结尾",
        [
            "通过这条演示链路可以看到，系统已经覆盖从设计知识检索、需求理解、模型生成、人工确认、预评审校验、变更影响分析到冲突合并入库的完整闭环。",
            "平台的价值不只是调用大模型，而是把大模型能力嵌入到 MBSE 的数据治理流程中：所有 AI 生成内容都可追溯、可校验、可审核、可合并，并最终沉淀为发布分支中的权威设计知识。",
            "后续可继续扩展更多专业工具，例如 GMAT 轨道仿真、Simulink 动态模型仿真，以及更多甲方自定义建模规范和审批流程。",
        ],
    )

    add_section_title(doc, "十一、录制检查清单")
    for item in [
        "登录页、角色切换、顶部系统状态标签均已录到。",
        "问答环节至少展示一个可追溯来源或模型元素定位。",
        "模型生成环节至少展示 SysML 文本和一张图形化视图。",
        "人在回路环节必须出现“预览/确认/采纳/修正”之一。",
        "校验环节必须出现问题报告和修复建议。",
        "变更影响环节必须出现影响拓扑或影响矩阵。",
        "合并入库环节必须出现冲突处理、合并按钮和审计日志。",
        "结尾用 20 秒总结完整闭环，不要只停留在单点功能。",
    ]:
        add_bullet(doc, item)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build_doc()
    print(out)
