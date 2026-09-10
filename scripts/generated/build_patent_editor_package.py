from pathlib import Path
import math
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\ai_mbse1\AI-MBSE-Demo")
OUT_DIR = ROOT / "docs" / "generated"
FIG_DIR = OUT_DIR / "patent_figures"
DOCX = OUT_DIR / "一种面向复杂系统工程模型的知识增强式生成校验与变更治理方法及系统_专利编辑材料包.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(80, 80, 80)
BLACK = RGBColor(0, 0, 0)


def find_font(size=32):
    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


FONT = find_font(34)
FONT_SMALL = find_font(28)
FONT_TINY = find_font(24)
FONT_TITLE = find_font(42)


def wrap_text(text, max_chars):
    lines = []
    for part in text.split("\n"):
        if len(part) <= max_chars:
            lines.append(part)
        else:
            lines.extend(textwrap.wrap(part, width=max_chars, break_long_words=False, replace_whitespace=False))
    return "\n".join(lines)


def box(draw, xy, text, font=FONT, max_chars=8, fill="white", width=3):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=10, outline="black", width=width, fill=fill)
    msg = wrap_text(text, max_chars)
    bbox = draw.multiline_textbbox((0, 0), msg, font=font, spacing=6, align="center")
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2), msg, fill="black", font=font, spacing=6, align="center")


def arrow(draw, start, end, width=4):
    draw.line([start, end], fill="black", width=width)
    sx, sy = start
    ex, ey = end
    ang = math.atan2(ey - sy, ex - sx)
    size = 16
    pts = [
        (ex, ey),
        (ex - size * math.cos(ang - math.pi / 6), ey - size * math.sin(ang - math.pi / 6)),
        (ex - size * math.cos(ang + math.pi / 6), ey - size * math.sin(ang + math.pi / 6)),
    ]
    draw.polygon(pts, fill="black")


def title(draw, text, w):
    bbox = draw.textbbox((0, 0), text, font=FONT_TITLE)
    draw.text(((w - (bbox[2] - bbox[0])) / 2, 30), text, fill="black", font=FONT_TITLE)


def save_fig(name, image):
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    path = FIG_DIR / name
    image.save(path)
    return path


def fig1():
    w, h = 2100, 700
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title(d, "图1  方法总体流程图", w)
    labels = [
        "多源工程\n数据获取",
        "工程对象\n抽取",
        "工程知识\n底座构建",
        "知识增强\n检索",
        "候选工程\n模型生成",
        "模型校验\n与修复",
        "人在回路\n审核",
        "沙箱变更\n影响分析",
        "合并入库\n形成权威模型",
    ]
    x0, y, bw, bh, gap = 55, 180, 195, 115, 30
    centers = []
    for i, lab in enumerate(labels):
        x = x0 + i * (bw + gap)
        box(d, (x, y, x + bw, y + bh), lab, FONT_SMALL, 6)
        centers.append((x + bw / 2, y + bh / 2))
        if i:
            arrow(d, (x - gap + 5, y + bh / 2), (x - 8, y + bh / 2), 3)
    d.rectangle((380, 410, 1720, 570), outline="black", width=3)
    d.text((430, 440), "闭环反馈：审核记录、校验结果、变更记录和仿真证据回写工程知识底座", font=FONT_SMALL, fill="black")
    arrow(d, (1540, 410), (760, 300), 3)
    arrow(d, (960, 410), (650, 300), 3)
    return save_fig("图1_方法总体流程图.png", img)


def fig2():
    w, h = 1700, 1100
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title(d, "图2  系统功能模块架构图", w)
    d.rounded_rectangle((60, 150, 1640, 860), radius=18, outline="black", width=4, fill="white")
    box(d, (620, 180, 1080, 275), "知识增强式工程模型治理系统", FONT_SMALL, 16)
    rows = [
        [
            ("数据接入模块", 140, 360),
            ("工程对象抽取模块", 500, 360),
            ("知识底座模块", 860, 360),
            ("知识增强检索模块", 1220, 360),
        ],
        [
            ("模型生成模块", 140, 560),
            ("模型校验模块", 500, 560),
            ("人在回路模块", 860, 560),
            ("变更治理模块", 1220, 560),
        ],
        [
            ("工具联动模块", 140, 760),
            ("合并入库模块", 500, 760),
            ("审计反馈模块", 860, 760),
            ("权威模型管理", 1220, 760),
        ],
    ]
    for row in rows:
        for lab, x, y in row:
            box(d, (x, y, x + 260, y + 105), lab, FONT_SMALL, 9)
    for row in rows:
        for i in range(len(row) - 1):
            x1, y1 = row[i][1] + 260, row[i][2] + 52
            x2, y2 = row[i + 1][1], row[i + 1][2] + 52
            arrow(d, (x1, y1), (x2, y2), 3)
    arrow(d, (1350, 465), (270, 560), 3)
    arrow(d, (1350, 665), (630, 760), 3)

    box(d, (100, 945, 500, 1030), "外部数据源\n文档/模型/仿真/评审", FONT_TINY, 16)
    box(d, (650, 945, 1050, 1030), "外部工程工具\n建模工具/仿真工具", FONT_TINY, 16)
    box(d, (1200, 945, 1600, 1030), "发布分支\n权威工程模型", FONT_TINY, 10)
    arrow(d, (300, 945), (270, 860), 3)
    arrow(d, (850, 945), (270, 865), 3)
    arrow(d, (1350, 865), (1400, 945), 3)
    return save_fig("图2_系统功能模块架构图.png", img)


def fig3():
    w, h = 1750, 1000
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title(d, "图3  工程知识底座结构示意图", w)
    box(d, (650, 130, 1100, 240), "工程知识底座", FONT, 10)
    top = [
        ("语义索引\n文档片段/历史问答/评审意见", 90, 390),
        ("工程知识图谱\n需求/模块/接口/参数/约束/验证项", 620, 390),
        ("工程规则库\n建模规则/校验规则/变更规则", 1150, 390),
    ]
    for lab, x, y in top:
        box(d, (x, y, x + 420, y + 150), lab, FONT_SMALL, 16)
        arrow(d, (875, 240), (x + 210, y), 3)
    objs = [
        ("需求", 280, 690),
        ("功能", 500, 760),
        ("模块", 730, 690),
        ("接口", 950, 760),
        ("参数", 1180, 690),
        ("约束", 1400, 760),
        ("验证项", 760, 875),
    ]
    for lab, x, y in objs:
        box(d, (x, y, x + 140, y + 70), lab, FONT_TINY, 4)
    for i in range(len(objs) - 1):
        arrow(d, (objs[i][1] + 140, objs[i][2] + 35), (objs[i + 1][1], objs[i + 1][2] + 35), 2)
    d.text((660, 600), "满足、派生、组成、连接、约束、验证、依赖、版本关系", font=FONT_SMALL, fill="black")
    return save_fig("图3_工程知识底座结构示意图.png", img)


def fig4():
    w, h = 1700, 980
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title(d, "图4  知识增强候选模型生成流程图", w)
    box(d, (80, 170, 390, 280), "用户建模任务\n或变更任务", FONT_SMALL, 12)
    box(d, (520, 170, 830, 280), "任务理解\n项目/角色/上下文", FONT_SMALL, 12)
    box(d, (960, 170, 1270, 280), "知识增强检索\n文本/图谱/规则", FONT_SMALL, 12)
    box(d, (1360, 170, 1630, 280), "生成上下文\n与来源证据", FONT_SMALL, 10)
    arrow(d, (390, 225), (520, 225), 3)
    arrow(d, (830, 225), (960, 225), 3)
    arrow(d, (1270, 225), (1360, 225), 3)
    box(d, (520, 430, 830, 540), "模型生成模块", FONT_SMALL, 10)
    box(d, (960, 430, 1270, 540), "候选工程模型", FONT_SMALL, 10)
    box(d, (960, 690, 1270, 820), "模型元素清单\n模型关系清单\n证据链", FONT_SMALL, 12)
    arrow(d, (1495, 280), (675, 430), 3)
    arrow(d, (830, 485), (960, 485), 3)
    arrow(d, (1115, 540), (1115, 690), 3)
    box(d, (80, 690, 390, 820), "可输出模型形式\nSysML/UML/AADL\n仿真接口/自定义模型", FONT_SMALL, 14)
    arrow(d, (960, 755), (390, 755), 3)
    return save_fig("图4_知识增强候选模型生成流程图.png", img)


def fig5():
    w, h = 1750, 1050
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title(d, "图5  候选模型校验与人在回路审核流程图", w)
    boxes = [
        ("候选工程模型", 110, 170),
        ("模型校验\n语法/语义/接口/约束", 520, 170),
        ("校验结果\n问题定位/风险等级", 930, 170),
        ("修复建议\n或模型修复补丁", 1320, 170),
        ("用户确认", 520, 470),
        ("人在回路审核\n采纳/修正/驳回/复核", 930, 470),
        ("候选状态更新", 1320, 470),
        ("反馈数据回写\n知识底座", 520, 760),
        ("审计记录", 930, 760),
        ("进入合并入库流程", 1320, 760),
    ]
    for lab, x, y in boxes:
        box(d, (x, y, x + 300, y + 115), lab, FONT_SMALL, 12)
    for a, b in [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (5, 6), (6, 9)]:
        sx, sy = boxes[a][1] + 300, boxes[a][2] + 57
        ex, ey = boxes[b][1], boxes[b][2] + 57
        if abs(sy - ey) < 10:
            arrow(d, (sx, sy), (ex, ey), 3)
    arrow(d, (670, 585), (670, 760), 3)
    arrow(d, (1080, 585), (1080, 760), 3)
    arrow(d, (820, 817), (930, 817), 3)
    arrow(d, (1230, 817), (1320, 817), 3)
    d.text((770, 360), "可自动修复的问题进入补丁确认；需工程判断的问题进入人工审核", font=FONT_TINY, fill="black")
    return save_fig("图5_候选模型校验与人在回路审核流程图.png", img)


def fig6():
    w, h = 1850, 1080
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title(d, "图6  沙箱变更影响分析与合并入库流程图", w)
    box(d, (120, 165, 430, 280), "变更请求\n参数/接口/模型元素", FONT_SMALL, 13)
    box(d, (580, 165, 890, 280), "沙箱环境\n模拟变更", FONT_SMALL, 10)
    box(d, (1040, 165, 1350, 280), "影响传播分析\n图谱关系/规则约束", FONT_SMALL, 13)
    box(d, (1460, 165, 1760, 280), "影响结果\n拓扑/矩阵/冲突", FONT_SMALL, 12)
    arrow(d, (430, 222), (580, 222), 3)
    arrow(d, (890, 222), (1040, 222), 3)
    arrow(d, (1350, 222), (1460, 222), 3)
    targets = [
        ("受影响需求", 210, 500),
        ("受影响模块", 500, 640),
        ("受影响接口", 790, 500),
        ("受影响参数", 1080, 640),
        ("受影响验证项", 1370, 500),
    ]
    for lab, x, y in targets:
        box(d, (x, y, x + 230, y + 90), lab, FONT_TINY, 8)
        arrow(d, (1195, 280), (x + 115, y), 2)
    box(d, (560, 870, 870, 970), "提交合并请求", FONT_SMALL, 10)
    box(d, (1010, 870, 1320, 970), "审核人处理冲突", FONT_SMALL, 10)
    box(d, (1460, 870, 1760, 970), "合并发布分支\n形成权威模型", FONT_SMALL, 12)
    arrow(d, (440, 545), (560, 920), 2)
    arrow(d, (870, 920), (1010, 920), 3)
    arrow(d, (1320, 920), (1460, 920), 3)
    return save_fig("图6_沙箱变更影响分析与合并入库流程图.png", img)


def fig7():
    w, h = 1600, 1050
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title(d, "图7  追溯关系示意图", w)
    center = (800, 520)
    box(d, (630, 460, 970, 580), "工程模型元素", FONT_SMALL, 10)
    nodes = [
        ("来源需求", 120, 170),
        ("来源文档片段", 630, 120),
        ("接口/参数/约束", 1130, 170),
        ("仿真验证证据", 1150, 760),
        ("审核记录", 630, 820),
        ("版本与审计日志", 120, 760),
    ]
    for lab, x, y in nodes:
        box(d, (x, y, x + 300, y + 105), lab, FONT_SMALL, 12)
        arrow(d, (x + 150, y + 105 if y < center[1] else y), center, 2)
    d.text((690, 650), "可追溯、可校验、可审核、可合并", font=FONT_SMALL, fill="black")
    return save_fig("图7_追溯关系示意图.png", img)


def set_run_font_docx(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    header = sec.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("专利编辑材料包")
    set_run_font_docx(r, 9, color=GRAY)
    footer = sec.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统")
    set_run_font_docx(r, 9, color=GRAY)
    return doc


def para(doc, text="", size=11, bold=False, color=None, before=0, after=6, align=None, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font_docx(r, size, bold, color)
    return p


def heading_doc(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12 if level == 2 else 8)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6 if level == 2 else 4)
    r = p.add_run(text)
    set_run_font_docx(r, 16 if level == 1 else 13 if level == 2 else 12, True, BLUE if level < 3 else DARK_BLUE)


def bullet_doc(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font_docx(r, 11)


def table_doc(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        p = table.cell(0, i).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font_docx(r, 10.5, True, DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font_docx(r, 10)
    para(doc, "", after=4)


def add_figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.4))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font_docx(r, 10.5, True, BLACK)


def build_doc(figs):
    doc = setup_doc()
    para(doc, "专利编辑材料包", 22, True, BLACK, after=4)
    para(doc, "一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统", 14, color=GRAY, after=12)

    heading_doc(doc, "一、基础信息", 1)
    table_doc(
        doc,
        ["项目", "内容"],
        [
            ("拟申请类型", "发明专利"),
            ("拟申请名称", "一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统"),
            ("技术领域", "模型驱动系统工程、工程模型数据治理、知识增强人工智能辅助设计。"),
            ("建议保护对象", "一种从工程数据获取、知识增强生成、模型校验、人在回路审核、沙箱变更分析到合并入库的通用工程模型治理方法及系统。"),
        ],
        [1600, 7760],
    )

    heading_doc(doc, "二、摘要建议稿", 1)
    para(doc, "本发明公开了一种面向复杂系统工程模型的知识增强式生成、校验与变更治理方法及系统，用于解决复杂系统建模门槛高、直接使用人工智能生成模型不可追溯且准确性不足、以及任务或模型变更时影响范围难以分析的问题。该方法获取复杂系统设计过程中的多源工程数据，抽取需求、功能、模块、接口、参数、约束和验证项等工程对象，并构建包括语义索引、工程知识图谱和工程规则库的工程知识底座；响应于建模任务或变更任务，基于工程知识底座进行知识增强检索，生成带有来源证据的候选工程模型；在候选工程模型入库前，对其进行模型校验并通过人在回路机制完成审核；当工程模型发生变更时，在沙箱环境中进行影响分析和冲突识别；经审核后，将候选工程模型或变更结果合并为权威工程模型，并生成审计记录。本发明能够提高复杂系统工程模型生成效率和入库质量，实现人工智能生成内容的可追溯、可校验、可审核和可治理。")

    heading_doc(doc, "三、现有技术及不足", 1)
    para(doc, "现有复杂系统工程建模通常依赖人工完成需求解析、功能分解、接口定义、模型构建、约束配置和验证关系维护。该过程要求建模人员同时具备领域知识、模型语言知识和工程工具使用经验，专业门槛较高、流程较为复杂，且不同人员的建模习惯和理解偏差会导致模型质量不稳定、模型维护成本较高。")
    para(doc, "虽然人工智能技术能够根据自然语言生成文本、代码或部分模型描述，但若直接使用通用人工智能生成工程模型，生成结果通常缺少工程知识、历史模型、接口规范和规则库约束，容易出现模型元素来源不可追溯、接口关系不完整、工程语义不一致、约束缺失或结果准确性不足等问题，难以直接作为权威工程模型入库。")
    para(doc, "此外，复杂系统中的需求、任务、参数或接口发生变更时，往往会影响多个模块、接口、约束、验证项、仿真任务和下游模型。现有方法通常依赖人工经验进行影响排查，难以全面识别直接影响和间接影响，也缺少在变更提交前进行沙箱式预演、冲突识别和可视化影响分析的机制。")
    para(doc, "因此，需要一种能够降低复杂系统工程建模门槛、约束人工智能生成过程、提供来源追溯和模型校验能力，并在模型或任务发生变更时自动分析影响范围的工程模型生成、校验与变更治理方法。")

    heading_doc(doc, "四、拟解决的技术问题", 1)
    for item in [
        "降低复杂系统工程建模对专业建模语言、建模工具和人工经验的依赖，提升需求到工程模型的转换效率。",
        "解决直接使用人工智能生成工程模型时存在的来源不可追溯、关系不完整、工程语义不一致和准确性不足的问题。",
        "解决候选工程模型入库前缺乏规则校验、修复建议、人工审核和状态治理的问题。",
        "解决任务、需求、参数或接口发生变更时，影响范围难以全面分析、间接影响容易遗漏、冲突难以及时识别的问题。",
        "解决AI生成模型、人工修正结果、仿真证据和发布分支权威模型之间缺乏闭环追溯和审计记录的问题。",
    ]:
        bullet_doc(doc, item)

    heading_doc(doc, "五、技术方案", 1)
    steps = [
        "S1，获取多源工程数据，包括需求文档、接口规范、历史工程模型、仿真结果、评审记录、版本记录、权限数据和用户反馈数据。",
        "S2，从多源工程数据中抽取工程对象，所述工程对象包括需求、功能、模块、接口、参数、约束和验证项。",
        "S3，基于工程对象构建工程知识底座，所述工程知识底座包括语义索引、工程知识图谱和工程规则库。",
        "S4，响应于建模任务或变更任务，基于工程知识底座进行知识增强检索，得到工程上下文和来源证据。",
        "S5，基于工程上下文生成候选工程模型，候选工程模型包括文本化模型描述、图形化模型视图、模型元素清单、模型关系清单和证据链。",
        "S6，对候选工程模型进行模型校验，校验内容包括语法、语义、接口完整性、属性类型、连接合法性、约束一致性和需求覆盖性。",
        "S7，通过人在回路机制接收用户对候选工程模型的采纳、修正、驳回、复核或审批操作。",
        "S8，针对工程模型变更，在沙箱环境中进行影响分析和冲突识别。",
        "S9，经审核后，将候选工程模型或变更结果合并为权威工程模型，并生成审计记录。",
    ]
    for item in steps:
        bullet_doc(doc, item)

    heading_doc(doc, "六、核心创新点", 1)
    table_doc(
        doc,
        ["创新点", "说明"],
        [
            ("知识增强生成", "利用语义索引、工程知识图谱和工程规则库共同形成生成上下文，使AI生成结果受工程知识约束。"),
            ("候选模型治理", "AI生成结果先处于候选状态，需经过校验和人工审核后才能转化为权威工程模型。"),
            ("校验与修复闭环", "在入库前执行模型语言、工程语义、接口、约束和需求覆盖等规则校验，并生成修复建议。"),
            ("沙箱变更预演", "变更先在沙箱环境中进行影响传播分析，不直接覆盖发布分支权威模型。"),
            ("证据链与审计", "需求、模型元素、仿真证据、审核记录和版本变更之间形成可追溯链路。"),
        ],
        [1900, 7460],
    )

    heading_doc(doc, "七、系统组成", 1)
    table_doc(
        doc,
        ["模块", "功能"],
        [
            ("数据接入模块", "获取需求文档、模型文件、接口规范、仿真结果、评审记录和历史版本数据。"),
            ("工程对象抽取模块", "抽取需求、功能、模块、接口、参数、约束和验证项等工程对象。"),
            ("知识底座模块", "维护语义索引、工程知识图谱和工程规则库。"),
            ("知识增强检索模块", "根据任务召回相关文本、模型元素、图谱关系和规则条目。"),
            ("模型生成模块", "生成候选工程模型、模型元素关系和图形化视图。"),
            ("模型校验模块", "执行语法、语义、接口、命名、属性、连接和约束校验。"),
            ("人在回路模块", "接收采纳、修正、驳回、复核和审批操作。"),
            ("变更治理模块", "执行沙箱影响分析、冲突识别和变更预演。"),
            ("合并入库模块", "将满足条件的候选模型或变更结果合并为权威工程模型。"),
            ("审计反馈模块", "记录审计信息并将反馈数据回写工程知识底座。"),
        ],
        [2100, 7260],
    )

    heading_doc(doc, "八、具体实施例", 1)
    heading_doc(doc, "实施例一：通用工程模型治理流程", 2)
    para(doc, "在一个实施例中，系统接入需求文档、接口规范、历史工程模型、仿真结果、评审记录、版本记录和用户反馈数据。系统对上述数据进行解析和抽取，形成需求、功能、模块、接口、参数、约束和验证项等工程对象。每个工程对象包括对象标识、对象类型、名称、描述、来源位置、生命周期状态和版本信息。")
    para(doc, "系统基于工程对象构建工程知识底座。语义索引用于召回相似文本和历史工程资料，工程知识图谱用于表示工程对象之间的满足、派生、组成、连接、约束、验证、依赖和版本关系，工程规则库用于存储模型语言规则、项目命名规则、接口连接规则、变更传播规则和入库规则。")

    heading_doc(doc, "实施例二：面向MBSE系统的具体应用", 2)
    para(doc, "在一个具体应用中，上述方法应用于AI赋能MBSE系统。系统接入需求文档、SysML模型、图数据库、向量数据库、工程建模工具同步结果、仿真工具结果和用户审核记录。系统基于知识增强检索召回相关需求、接口、模块和约束，并生成SysML文本模型、需求图、用例图、内部模块图或接口关系图。")
    para(doc, "生成的SysML模型在入库前以候选模型形式展示，设计人员可以进行采纳、修正或驳回。系统根据SysML语法、项目命名规范、接口完整性、属性类型、连接合法性和需求覆盖性进行预评审与校验，并对可修复问题给出修复建议。")
    para(doc, "当设计人员提出参数或接口变更时，例如修改链路带宽、接口容量或载荷参数，系统在沙箱环境中进行影响传播分析，输出受影响需求、模块、接口、验证项和仿真任务。审核通过后，系统通过分支合并将候选模型或变更结果合并为发布分支中的权威模型元素。")

    heading_doc(doc, "实施例三：仿真证据关联", 2)
    para(doc, "对于需要仿真验证的需求，系统可以调用外部仿真工具生成验证证据。例如，在航天系统设计场景中，轨道覆盖类需求可以关联轨道仿真结果，动态响应类需求可以关联动态仿真结果。仿真结果与需求、参数、模型元素和校验项建立追溯关系，用于支持模型校验、工程评审和后续变更分析。")

    heading_doc(doc, "九、可替代实施方式", 1)
    for item in [
        "语义索引可以由向量数据库、全文检索引擎、关键词索引或其组合实现。",
        "工程知识图谱可以由图数据库、关系型数据库、文档数据库或内存图结构实现。",
        "候选工程模型可以由大语言模型、规则模板、模型转换器或多模型协同机制生成。",
        "工程规则库可以包括标准建模语言规范、企业建模规范、接口控制规则、仿真约束规则或项目自定义规则。",
        "工程模型不限于SysML，也可以包括UML、AADL、Simulink接口模型、数字孪生模型或企业自定义工程模型。",
        "本发明不限于航天系统，也可以应用于航空、船舶、车辆、能源、工业装备、智能制造或其他复杂系统工程设计场景。",
    ]:
        bullet_doc(doc, item)

    heading_doc(doc, "十、建议权利要求布局", 1)
    table_doc(
        doc,
        ["类型", "建议保护内容"],
        [
            ("方法主权利要求", "保护多源数据获取、工程对象抽取、知识底座构建、知识增强检索、候选模型生成、模型校验、人在回路、沙箱变更分析和合并入库的完整闭环。"),
            ("方法从属权利要求", "分别限定多源工程数据类型、知识图谱关系、知识增强检索方式、模型校验内容、修复补丁、审核操作、沙箱影响分析和工具联动。"),
            ("系统权利要求", "按数据接入、对象抽取、知识底座、检索、生成、校验、人在回路、变更治理、合并入库和审计反馈模块保护。"),
            ("设备和介质权利要求", "补充电子设备和计算机可读存储介质保护形式。"),
        ],
        [2200, 7160],
    )

    heading_doc(doc, "十一、附图", 1)
    captions = [
        "图1  方法总体流程图",
        "图2  系统功能模块架构图",
        "图3  工程知识底座结构示意图",
        "图4  知识增强候选模型生成流程图",
        "图5  候选模型校验与人在回路审核流程图",
        "图6  沙箱变更影响分析与合并入库流程图",
        "图7  追溯关系示意图",
    ]
    for path, caption in zip(figs, captions):
        add_figure(doc, path, caption)

    heading_doc(doc, "十二、附图说明", 1)
    for item in captions:
        bullet_doc(doc, item + "。")

    doc.save(DOCX)


def build():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    figures = [fig1(), fig2(), fig3(), fig4(), fig5(), fig6(), fig7()]
    build_doc(figures)
    print(DOCX)
    for fig in figures:
        print(fig)


if __name__ == "__main__":
    build()
